"""
Script to generate a fresh, simple, and comprehensive End-to-End Guide
for SwinkPay Reconciliation, Automated Fixes, XCD Input Files, and Invoicing.
Generates:
- Desktop: SwinkPay_Reconciliation_End_To_End_Guide.docx
- Project: SwinkPay_Reconciliation_End_To_End_Guide.docx
- Project: SwinkPay_Reconciliation_End_To_End_Guide.md
"""
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_TITLE = "SwinkPay Daily Reconciliation & Operations Guide"
DOC_SUBTITLE = "Simple End-to-End Guide: Recon, Automated Fixes, XCD Input Files & Invoice Generation"
DOC_VERSION = "Version 2.0 (Simplified)"
DOC_DATE = "September 2026"
DOC_AUTHOR = "SwinkPay Fintech Operations & Automation Team"

# Styling colors
COLOR_PRIMARY = RGBColor(30, 58, 138)   # #1e3a8a
COLOR_ACCENT = RGBColor(79, 70, 229)    # #4f46e5
COLOR_TEXT = RGBColor(30, 41, 59)       # #1e293b
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748b
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_SUCCESS = RGBColor(16, 185, 129)  # #10b981

HEX_HEADER_BG = "1E3A8A"
HEX_ROW_ALT = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F0Fdf4"      # light green / light blue
HEX_CALLOUT_BORDER = "10B981"  # green
HEX_BOX_BG = "EEF2FF"          # soft indigo
HEX_BOX_BORDER = "4F46E5"
HEX_CODE_BG = "F1F5F9"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout(doc, title, text, bg_hex=HEX_CALLOUT_BG, border_hex=HEX_CALLOUT_BORDER):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_margins(cell, top=140, bottom=140, left=180, right=160)
    set_cell_shading(cell, bg_hex)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_PRIMARY
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = COLOR_TEXT
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def add_diagram_box(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_shading(cell, HEX_CODE_BG)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def create_styled_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=HEX_BORDER)
    
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = COLOR_WHITE
        if col_widths and i < len(col_widths):
            cell.width = Inches(col_widths[i])
            
    for row_idx, r_data in enumerate(rows_data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = HEX_ROW_ALT if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, cell_value in enumerate(r_data):
            cell = row_cells[col_idx]
            if bg_color != "FFFFFF":
                set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(cell_value))
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEXT
            if col_widths and col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return table


def generate_simple_docx(output_path: str):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_TEXT
    
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(DOC_TITLE)
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run(DOC_SUBTITLE)
    run_sub.font.size = Pt(11.5)
    run_sub.font.color.rgb = COLOR_MUTED
    
    add_callout(
        doc,
        "What is this document?",
        "This guide explains SwinkPay's daily reconciliation and settlement process in simple, practical terms. "
        "It covers what happens every morning, what we do when numbers match ('All Fine'), how we fix the 3 common issues "
        "(Pull, Sync, Network Change), and how final settlement files (XCD) and merchant invoices are generated."
    )
    
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_ACCENT
        return p

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT
        return p

    # SECTION 1
    add_h1("1. The Big Picture: What Are We Doing & Why?")
    add_p(
        "Every day, thousands of customers scan SwinkPay static QR codes at merchant outlets (like Coffee Day, etc.) "
        "and make UPI payments. The money flows through three Payment Gateways (Cashfree, Easebuzz, Airtel Bank)."
    )
    add_p(
        "At the end of every day, SwinkPay has three main goals:\n"
        "1. Audit & Match: Ensure every rupee collected by the Payment Gateways is in SwinkPay's system (Zero Loss).\n"
        "2. Merchant Settlement: Pay each merchant their money (Gross amount collected minus SwinkPay's fees).\n"
        "3. Invoicing: Generate tax invoices for SwinkPay's commission/MDR fees + GST charged to the merchant."
    )
    add_p(
        "To make this happen, every morning at T+1, the Operations team runs Daily Reconciliation."
    )
    
    # SECTION 2
    add_h1("2. Flow 1: When Everything is 'ALL FINE' (The Happy Path)")
    add_p(
        "When all files match perfectly on the morning run (or after fixing any discrepancies), the difference is exactly 0.00. "
        "Here is what happens step-by-step:"
    )
    
    flow1_diagram = (
        "[Step 1: Upload Reports] ──> [Step 2: Run Recon] ──> [Difference is 0.00 (ALL CLEAR)]\n"
        "                                                                │\n"
        "            ┌───────────────────────────────────────────────────┼─────────────────────────────────┐\n"
        "            ▼                                                   ▼                                 ▼\n"
        "[Step 3: Download XCD Input Files]                    [Step 4: Invoice Generation]      [Step 5: Daily Recon Email]\n"
        "  • XCD Cashfree (.xlsx)                                • CMS calculates exact MDR fees   • Copy summary from dashboard\n"
        "  • XCD Easebuzz (.xlsx)                                • Generates GST Tax Invoices      • Send to Finance with XCD\n"
        "  • XCD Airtel (.xlsx)                                    for each merchant outlet          attachments\n"
        "  └──> Sent to Bank to credit merchants' accounts"
    )
    add_diagram_box(doc, flow1_diagram)
    
    add_h2("What are XCD Input Files?")
    add_p(
        "XCD Input files are the official settlement instruction workbooks. There is one file per payment gateway:\n"
        "• XCD Input file as on YYYY-MM-DD (Cf).xlsx\n"
        "• XCD Input file as on YYYY-MM-DD (EB).xlsx\n"
        "• XCD Input file as on YYYY-MM-DD (Airtel).xlsx\n"
        "\n"
        "What happens with them? These files are sent to the nodal bank / settlement account. "
        "The bank uses them to credit the Net Settlement Amount into each merchant's bank account."
    )
    
    add_h2("What is Invoice Generation?")
    add_p(
        "Because reconciliation has verified every transaction, SwinkPay can legally and accurately bill the merchant:\n"
        "• Gross Amount Collected (e.g. ₹100.00)\n"
        "• Service Charge / MDR Fee (e.g. ₹1.65)\n"
        "• GST @ 18% (e.g. ₹0.30)\n"
        "• Net Paid to Merchant (₹98.05)\n"
        "\n"
        "CMS generates a Tax Invoice for the merchant covering the MDR fee and GST."
    )
    
    add_h2("Daily Recon Email")
    add_p(
        "Operations clicks 'Copy Formatted Text' on the Ops_Auto dashboard. It generates a clean executive email "
        "with outlet count, Cashfree pulled count, and partner settlement table. Operations pastes this into email, "
        "attaches the XCD workbooks, and sends it to Finance."
    )

    # SECTION 3
    add_h1("3. Flow 2: What If There Are Issues? (The 3 Simple Fixes)")
    add_p(
        "If the initial run does NOT balance to 0.00, it is ALWAYS due to one of three common operational issues. "
        "Ops_Auto makes fixing all three simple and automated:"
    )
    
    issues_table = [
        ["1. Missing in CMS", "Customer paid on QR, Gateway captured money, but CMS missed it.", "Click 'Pull' on dashboard. Ops_Auto looks up Terminal ID & pulls txn via API.", "Backend adds transaction into CMS."],
        ["2. Missing in SMMS", "Transaction is in CMS, but SMMS sync status is 'false' (un-synced).", "Download 'Sync_Transactions.xlsx' from dashboard badge.", "Upload file to CMS/SMMS portal under Bulk Sync."],
        ["3. Network Mismatch", "CMS recorded wrong network (e.g. RUPAY/WA instead of UPI Credit/PPI).", "Download 'Change_Network.xlsx' (exactly 91 genuine changes).", "Upload file to CMS portal under Bulk Network Update."]
    ]
    create_styled_table(doc, ["Issue", "What Happened?", "How We Fix It in Ops_Auto", "Result / Next Step"], issues_table, [1.4, 1.8, 1.8, 1.5])
    
    add_h2("Deep-Dive on Fix 1: Pulling Missing Transactions (No Postman Required)")
    add_p(
        "When money is in the gateway report but missing in CMS, we must 'Pull' it so CMS creates a record:\n"
        "• Easebuzz & Airtel: The Terminal ID is already written in the PG report.\n"
        "• Cashfree: The Terminal ID is embedded in the Order ID (e.g. '330595-4860-AXI...'). "
        "The engine copies the middle code ('4860'), searches your uploaded Terminal File, finds the matching Terminal ID, "
        "and prepares the pull request.\n"
        "• Operations clicks 'Pull Missing Transactions' directly on the dashboard. The system pulls all missing records "
        "directly into CMS without opening Postman!"
    )
    
    add_h2("Deep-Dive on Fix 2: Syncing Transactions to SMMS")
    add_p(
        "When transactions are in CMS but missing in SMMS, SMMS cannot calculate merchant settlements. "
        "Ops_Auto generates 'Sync_Transactions_YYYY-MM-DD.xlsx'. Operations downloads this file and uploads it into "
        "the CMS/SMMS Admin Portal. The backend runs a sync job and marks them synced (status = true)."
    )
    
    add_h2("Deep-Dive on Fix 3: Changing Network (The 91 Transactions)")
    add_p(
        "CMS sometimes mislabels Rupay Credit Cards on UPI as standard card 'RUPAY', or Wallets as 'WA'.\n"
        "• Standard UPI ('UPI' vs 'UPI_OFFLINE_STATIC') is normal static QR and is NOT a change.\n"
        "• The engine isolates ONLY the 91 real misclassifications: 77 RUPAY (Credit Card UPI) + 14 WA (Wallet UPI).\n"
        "• Ops_Auto generates 'Change_Network_YYYY-MM-DD.xlsx' (Columns: SwinkPay Txn ID, Old Network, New Network).\n"
        "• Operations uploads this file to the CMS Bulk Network Update screen. CMS updates the network field."
    )
    
    # SECTION 4
    add_h1("4. The Complete Everyday Cycle (In 1 Picture)")
    
    loop_diagram = (
        "                    Start Day (Morning T+1)\n"
        "                              │\n"
        "                              ▼\n"
        "                    Run Initial Recon\n"
        "                              │\n"
        "               Are there any discrepancies?\n"
        "                              │\n"
        "           ┌──────────────────┴──────────────────┐\n"
        "          YES                                   NO\n"
        "           │                                     │\n"
        "Fix the 3 things:                                │\n"
        "1. Click 'Pull' for missing PG transactions      │\n"
        "2. Upload Sync file to CMS/SMMS                  │\n"
        "3. Upload Network Change file to CMS             │\n"
        "           │                                     │\n"
        "           ▼                                     │\n"
        "Re-download fresh CMS & SMMS                     │\n"
        "           │                                     │\n"
        "           ▼                                     │\n"
        "   Re-run Recon ─────────────────────────────────┘\n"
        "                              │\n"
        "                              ▼\n"
        "                    Difference is 0.00!\n"
        "                     (ALL IS CLEAR)\n"
        "                              │\n"
        "    ┌─────────────────────────┼─────────────────────────┐\n"
        "    ▼                         ▼                         ▼\n"
        "Download XCD            Generate Invoices          Send Email\n"
        "Input Files             for Merchant Fees          to Finance\n"
        "    │                         │                         │\n"
        "    ▼                         ▼                         ▼\n"
        "Money settled into       Invoices recorded         Daily Recon Closed\n"
        "merchant bank accounts    in billing system           Successfully!\n"
    )
    add_diagram_box(doc, loop_diagram)
    
    # SECTION 5
    add_h1("5. Everyday Operations Checklist")
    add_p("Follow this simple 7-step checklist every morning:")
    
    chk_data = [
        ["1", "Ingest Reports", "Upload CMS, SMMS, Cashfree, Easebuzz, Airtel reports", "[  ] Done"],
        ["2", "Pull Missing PG", "If any txns in PG but not CMS, click 'Pull' on dashboard", "[  ] Done"],
        ["3", "Upload Sync File", "Download Sync_Transactions.xlsx and upload to CMS/SMMS portal", "[  ] Done"],
        ["4", "Upload Network File", "Download Change_Network.xlsx (91 rows) and upload to CMS portal", "[  ] Done"],
        ["5", "Re-run Recon", "Re-run recon on refreshed reports -> Verify difference is 0.00", "[  ] Done"],
        ["6", "Download XCD Files", "Download partner XCD workbooks for settlement banking", "[  ] Done"],
        ["7", "Invoice & Email", "Invoices generated in CMS; copy email summary and send to Finance", "[  ] Done"]
    ]
    create_styled_table(doc, ["Step", "Action", "Description", "Status"], chk_data, [0.6, 1.8, 3.2, 0.9])
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"[SUCCESS] Saved Word doc to: {output_path}")


def generate_simple_md(output_path: str):
    content = f"""# {DOC_TITLE}
## {DOC_SUBTITLE}

**Version**: {DOC_VERSION}  
**Date**: {DOC_DATE}  
**Author**: {DOC_AUTHOR}  

---

## 1. The Big Picture: What Are We Doing & Why?

Every day, thousands of customers scan SwinkPay static QR codes at merchant outlets (like Coffee Day, etc.) and make UPI payments. The money flows through three Payment Gateways (**Cashfree, Easebuzz, Airtel Bank**).

At the end of every day, SwinkPay has three main goals:
1. **Audit & Match**: Ensure every rupee collected by the Payment Gateways is accounted for in SwinkPay's system (Zero Loss).
2. **Merchant Settlement**: Pay each merchant their money (Gross amount collected minus SwinkPay's fees).
3. **Invoicing**: Generate tax invoices for SwinkPay's commission/MDR fees + GST charged to the merchant.

To make this happen, every morning at **T+1**, the Operations team runs **Daily Reconciliation**.

---

## 2. Flow 1: When Everything is "ALL FINE" (The Happy Path)

When all files match perfectly on the morning run (or after fixing any discrepancies), the difference is exactly **0.00**.

```text
[Step 1: Upload Reports] ──> [Step 2: Run Recon] ──> [Difference is 0.00 (ALL CLEAR)]
                                                                │
            ┌───────────────────────────────────────────────────┼─────────────────────────────────┐
            ▼                                                   ▼                                 ▼
[Step 3: Download XCD Input Files]                    [Step 4: Invoice Generation]      [Step 5: Daily Recon Email]
  • XCD Cashfree (.xlsx)                                • CMS calculates exact MDR fees   • Copy summary from dashboard
  • XCD Easebuzz (.xlsx)                                • Generates GST Tax Invoices      • Send to Finance with XCD
  • XCD Airtel (.xlsx)                                    for each merchant outlet          attachments
  └──> Sent to Bank to credit merchants' accounts
```

### What are XCD Input Files?
XCD Input files are the official settlement instruction workbooks. There is one file per payment gateway:
- `XCD Input file as on YYYY-MM-DD (Cf).xlsx`
- `XCD Input file as on YYYY-MM-DD (EB).xlsx`
- `XCD Input file as on YYYY-MM-DD (Airtel).xlsx`

**What happens with them?**  
These files are sent to the nodal bank / settlement account. The bank uses them to credit the **Net Settlement Amount** into each merchant's bank account.

### What is Invoice Generation?
Because reconciliation has verified every transaction, SwinkPay can legally and accurately bill the merchant:
- **Gross Amount Collected** (e.g. ₹100.00)
- **Service Charge / MDR Fee** (e.g. ₹1.65)
- **GST @ 18%** (e.g. ₹0.30)
- **Net Paid to Merchant** (₹98.05)

CMS generates a **Tax Invoice** for the merchant covering the MDR fee and GST.

### Daily Recon Email
Operations clicks **"Copy Formatted Text"** on the Ops_Auto dashboard. It generates a clean executive email with outlet count, Cashfree pulled count, and partner settlement table. Operations pastes this into email, attaches the XCD workbooks, and sends it to Finance.

---

## 3. Flow 2: What If There Are Issues? (The 3 Simple Fixes)

If the initial run does NOT balance to 0.00, it is ALWAYS due to one of three common operational issues. Ops_Auto makes fixing all three simple and automated:

| Issue | What Happened? | How We Fix It in Ops_Auto | Result / Next Step |
| :--- | :--- | :--- | :--- |
| **1. Missing in CMS** | Customer paid on QR, Gateway captured money, but CMS missed it. | Click **"Pull"** on dashboard. Ops_Auto looks up Terminal ID & pulls txn via API. | Backend adds transaction into CMS. |
| **2. Missing in SMMS** | Transaction is in CMS, but SMMS sync status is `false` (un-synced). | Download **`Sync_Transactions.xlsx`** from dashboard badge. | Upload file to CMS/SMMS portal under Bulk Sync. |
| **3. Network Mismatch** | CMS recorded wrong network (e.g. RUPAY/WA instead of UPI Credit/PPI). | Download **`Change_Network.xlsx`** (exactly 91 genuine changes). | Upload file to CMS portal under Bulk Network Update. |

### Deep-Dive on Fix 1: Pulling Missing Transactions (No Postman Required)
When money is in the gateway report but missing in CMS, we must "Pull" it so CMS creates a record:
- **Easebuzz & Airtel**: The Terminal ID is already written in the PG report.
- **Cashfree**: The Terminal ID is embedded in the Order ID (e.g. `330595-4860-AXI...`). The engine copies the middle code (`4860`), searches your uploaded Terminal File, finds the matching Terminal ID, and prepares the pull request.
- Operations clicks **"Pull Missing Transactions"** directly on the dashboard. The system pulls all missing records directly into CMS without opening Postman!

### Deep-Dive on Fix 2: Syncing Transactions to SMMS
When transactions are in CMS but missing in SMMS, SMMS cannot calculate merchant settlements. Ops_Auto generates `Sync_Transactions_YYYY-MM-DD.xlsx`. Operations downloads this file and uploads it into the CMS/SMMS Admin Portal. The backend runs a sync job and marks them synced (`status = true`).

### Deep-Dive on Fix 3: Changing Network (The 91 Transactions)
CMS sometimes mislabels Rupay Credit Cards on UPI as standard card `RUPAY`, or Wallets as `WA`.
- Standard UPI (`UPI` vs `UPI_OFFLINE_STATIC`) is normal static QR and is **NOT** a change.
- The engine isolates **ONLY the 91 real misclassifications**: 77 RUPAY (Credit Card UPI) + 14 WA (Wallet UPI).
- Ops_Auto generates `Change_Network_YYYY-MM-DD.xlsx` (Columns: `SwinkPay Transaction ID`, `Old Network`, `New Network`).
- Operations uploads this file to the CMS Bulk Network Update screen. CMS updates the network field.

---

## 4. The Complete Everyday Cycle (In 1 Picture)

```text
                    Start Day (Morning T+1)
                              │
                              ▼
                    Run Initial Recon
                              │
               Are there any discrepancies?
                              │
           ┌──────────────────┴──────────────────┐
          YES                                   NO
           │                                     │
Fix the 3 things:                                │
1. Click 'Pull' for missing PG transactions      │
2. Upload Sync file to CMS/SMMS                  │
3. Upload Network Change file to CMS             │
           │                                     │
           ▼                                     │
Re-download fresh CMS & SMMS                     │
           │                                     │
           ▼                                     │
   Re-run Recon ─────────────────────────────────┘
                              │
                              ▼
                    Difference is 0.00!
                     (ALL IS CLEAR)
                              │
    ┌─────────────────────────┼─────────────────────────┐
    ▼                         ▼                         ▼
Download XCD            Generate Invoices          Send Email
Input Files             for Merchant Fees          to Finance
    │                         │                         │
    ▼                         ▼                         ▼
Money settled into       Invoices recorded         Daily Recon Closed
merchant bank accounts    in billing system           Successfully!
```

---

## 5. Everyday Operations Checklist

| Step | Action | Description | Status |
| :---: | :--- | :--- | :---: |
| **1** | **Ingest Reports** | Upload CMS, SMMS, Cashfree, Easebuzz, Airtel reports | `[ ]` |
| **2** | **Pull Missing PG** | If any txns in PG but not CMS, click 'Pull' on dashboard | `[ ]` |
| **3** | **Upload Sync File** | Download `Sync_Transactions.xlsx` and upload to CMS/SMMS portal | `[ ]` |
| **4** | **Upload Network File** | Download `Change_Network.xlsx` (91 rows) and upload to CMS portal | `[ ]` |
| **5** | **Re-run Recon** | Re-run recon on refreshed reports -> Verify difference is 0.00 | `[ ]` |
| **6** | **Download XCD Files** | Download partner XCD workbooks for settlement banking | `[ ]` |
| **7** | **Invoice & Email** | Invoices generated in CMS; copy email summary and send to Finance | `[ ]` |
"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"[SUCCESS] Saved Markdown to: {output_path}")


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(base_dir, "SwinkPay_Reconciliation_End_To_End_Guide.docx")
    md_path = os.path.join(base_dir, "SwinkPay_Reconciliation_End_To_End_Guide.md")
    
    print("Generating Simplified Word Document...")
    generate_simple_docx(docx_path)
    
    print("Generating Simplified Markdown Document...")
    generate_simple_md(md_path)
    
    # Copy to Desktop
    desktop_dir = "C:\\Users\\HP\\OneDrive - SwinkPay Fintech Pvt Ltd\\Desktop"
    if os.path.exists(desktop_dir):
        desktop_docx = os.path.join(desktop_dir, "SwinkPay_Reconciliation_End_To_End_Guide.docx")
        shutil.copyfile(docx_path, desktop_docx)
        print(f"[SUCCESS] Copied to Desktop: {desktop_docx}")
