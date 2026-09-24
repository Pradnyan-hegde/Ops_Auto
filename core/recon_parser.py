"""
Reconciliation Workbook Parser.
Parses an uploaded completed 'Reconciliation_YYYY-MM-DD.xlsx' workbook:
- Reads matched partner sheets: 'CMS_CF_Matched', 'CMS_EB_Matched', 'CMS_Air_Matched'
- Reads 'CMS' and 'Cashfree' to detect network discrepancies
- Reads 'CMS_Not_in_SMMS' or CMS sync status to detect SMMS sync requirements
- Computes distinct transaction dates and counts for multi-day date filtering
- Enables instant XCD downloads without re-running raw reconciliation
"""
import os
from typing import Dict, Any, List, Optional
import openpyxl

from .normalizer import clean_key, clean_amount
from .network_builder import detect_network_changes
from .xcd_builder import get_distinct_dates_with_counts, extract_record_date, extract_record_amounts


def _read_sheet_records(ws) -> List[Dict[str, Any]]:
    """Helper to convert an openpyxl worksheet into a list of row dicts."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return []

    headers = [str(c or "").strip() for c in rows[0]]
    records = []
    for r in rows[1:]:
        if not any(r):
            continue
        rec = {}
        for idx, h in enumerate(headers):
            if h:
                rec[h] = r[idx] if idx < len(r) else None
        records.append(rec)
    return records


def parse_reconciliation_workbook(file_path: str) -> Dict[str, Any]:
    """
    Parses a completed Reconciliation workbook (.xlsx).
    Returns a comprehensive dictionary with matched records, dates, network changes, and stats.
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames

    # Read matched sheets
    cms_cf_matched = []
    cms_eb_matched = []
    cms_air_matched = []

    if "CMS_CF_Matched" in sheet_names:
        cms_cf_matched = _read_sheet_records(wb["CMS_CF_Matched"])
    if "CMS_EB_Matched" in sheet_names:
        cms_eb_matched = _read_sheet_records(wb["CMS_EB_Matched"])
    if "CMS_Air_Matched" in sheet_names:
        cms_air_matched = _read_sheet_records(wb["CMS_Air_Matched"])

    # Read SMMS if present to build net amount lookup
    smms_net_map: Dict[str, float] = {}
    if "SMMS" in sheet_names:
        smms_records = _read_sheet_records(wb["SMMS"])
        for r in smms_records:
            net_val = clean_amount(r.get("Net Amount"))
            if net_val is not None:
                sp_id = str(r.get("SwinkPay Txn ID") or "").strip()
                if sp_id:
                    smms_net_map[sp_id] = net_val
                rrn = str(r.get("RRN/UTR") or r.get("RRN") or "").strip()
                if rrn:
                    smms_net_map[rrn] = net_val

    # Read raw CMS and Cashfree for network verification if present
    cms_records = []
    cf_records = []
    if "CMS" in sheet_names:
        cms_records = _read_sheet_records(wb["CMS"])
    if "Cashfree" in sheet_names:
        cf_records = _read_sheet_records(wb["Cashfree"])

    # Detect SMMS sync requirements
    cms_not_in_smms = []
    if "CMS_Not_in_SMMS" in sheet_names:
        cms_not_in_smms = _read_sheet_records(wb["CMS_Not_in_SMMS"])
    elif cms_records:
        # Check SMMS Sync Status in CMS rows
        for r in cms_records:
            sync_st = str(r.get("SMMS Sync Status") or "").strip().lower()
            if sync_st in ("false", "0", "no"):
                cms_not_in_smms.append(r)

    # Detect network changes
    network_changes = detect_network_changes(
        matched_cf_records=cms_cf_matched,
        cms_records=cms_records,
        cf_records=cf_records,
        cms_not_in_smms=cms_not_in_smms
    )

    # Compute dates and transaction counts with gross and net amounts
    dates_with_counts = get_distinct_dates_with_counts(
        cms_cf_matched,
        cms_eb_matched,
        cms_air_matched,
        net_lookup=smms_net_map
    )

    # Compute stats
    def calc_stats(recs):
        cnt = len(recs)
        gross_tot = 0.0
        net_tot = 0.0
        for r in recs:
            g, n = extract_record_amounts(r, smms_net_map)
            gross_tot += g
            net_tot += n
        return cnt, round(gross_tot, 2), round(net_tot, 2)

    cf_cnt, cf_gross, cf_net = calc_stats(cms_cf_matched)
    eb_cnt, eb_gross, eb_net = calc_stats(cms_eb_matched)
    air_cnt, air_gross, air_net = calc_stats(cms_air_matched)

    total_cnt = cf_cnt + eb_cnt + air_cnt
    total_gross = round(cf_gross + eb_gross + air_gross, 2)
    total_net = round(cf_net + eb_net + air_net, 2)

    # Compute distinct active outlets from Merchant MMS Terminal ID (pivot in CMS)
    distinct_terminals = set()
    source_outlet_recs = cms_records if cms_records else (cms_cf_matched + cms_eb_matched + cms_air_matched)
    for r in source_outlet_recs:
        tid = str(r.get("Merchant MMS Terminal ID") or "").strip()
        if tid and tid.lower() not in ("nan", "none", ""):
            distinct_terminals.add(tid)
    outlet_count = len(distinct_terminals)

    return {
        "cms_cf_matched": cms_cf_matched,
        "cms_eb_matched": cms_eb_matched,
        "cms_air_matched": cms_air_matched,
        "cms_not_in_smms": cms_not_in_smms,
        "network_changes": network_changes,
        "outlet_count": outlet_count,
        "dates": dates_with_counts,
        "stats": {
            "cashfree": {"count": cf_cnt, "gross_amount": cf_gross, "net_amount": cf_net},
            "easebuzz": {"count": eb_cnt, "gross_amount": eb_gross, "net_amount": eb_net},
            "airtel": {"count": air_cnt, "gross_amount": air_gross, "net_amount": air_net},
            "total": {"count": total_cnt, "gross_amount": total_gross, "net_amount": total_net}
        }
    }
