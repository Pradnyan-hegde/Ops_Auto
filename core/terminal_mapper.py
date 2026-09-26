"""
Terminal Reference Mapping Module.
Parses Terminal Report Excel files (TID_FILE) containing:
- TERMINAL ID
- MMS TERMINAL ID
- Partner Ref ID
- VPA

Maps Cashfree Order ID middle numbers (e.g. '4860' from '330595-4860-...')
to their authoritative MMS Terminal ID (e.g. 'X47VVH' / 'X1JYYO').
Persists mappings to data/terminal_mappings.json for instant lookups across sessions.
"""
import os
import re
import json
from typing import Dict, Any, Optional, Tuple, List
from datetime import datetime
import openpyxl

from .normalizer import clean_key


DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
MAPPINGS_FILE = os.path.join(DATA_DIR, "terminal_mappings.json")

# In-memory cached mappings
_CACHE: Optional[Dict[str, Any]] = None


def _normalize_header(val: Any) -> str:
    """Normalizes header text for robust comparison."""
    if val is None:
        return ""
    s = str(val).strip().lower()
    s = re.sub(r'[\s_\-/]+', ' ', s)
    return s.strip()


def parse_terminal_report(file_path: str) -> Dict[str, Any]:
    """
    Parses a Terminal Report workbook (.xlsx or .xls).
    Finds the header row containing 'TERMINAL ID', 'MMS TERMINAL ID', and 'Partner Ref ID'.
    Returns structured mappings indexed by Partner Ref ID, VPA, and Terminal ID.
    """
    import io
    wb = None
    sheet_rows = []

    # Read binary bytes to check OOXML vs legacy BIFF8
    with open(file_path, "rb") as f:
        file_bytes = f.read()

    magic = file_bytes[:8]
    if magic.startswith(b"PK\x03\x04"):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    elif magic.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=file_bytes)
            target_sheet = None
            for sname in book.sheet_names():
                if "tid" in sname.lower() or "terminal" in sname.lower():
                    target_sheet = book.sheet_by_name(sname)
                    break
            if not target_sheet:
                target_sheet = book.sheet_by_index(0)
            sheet_rows = [target_sheet.row_values(r) for r in range(target_sheet.nrows)]
        except Exception as e:
            raise RuntimeError(f"Could not read Excel .xls file {file_path}: {e}")
    else:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

    if wb is not None:
        sheet_name = None
        for name in wb.sheetnames:
            if "tid" in name.lower() or "terminal" in name.lower():
                sheet_name = name
                break
        ws = wb[sheet_name] if sheet_name else wb.active
        sheet_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

    header_row_idx = None
    col_map = {}

    # Scan rows to locate header row (up to 2000 rows)
    for r_idx, row in enumerate(sheet_rows[:2000]):
        norm_cells = [_normalize_header(c) for c in row]
        norm_dict = {norm_cells[i]: i for i in range(len(norm_cells)) if norm_cells[i]}

        # Check for key columns
        has_mms = any("mms terminal id" in k or "mms terminal" in k or "merchant mms terminal" in k for k in norm_dict)
        has_tid = any(k == "terminal id" or k == "terminalid" for k in norm_dict)
        has_pref = any("partner ref id" in k or "partner ref" in k or "ref id" in k for k in norm_dict)
        has_vpa = any("vpa" in k for k in norm_dict)

        if has_mms and (has_tid or has_pref or has_vpa):
            header_row_idx = r_idx
            # Map canonical column names to indices
            for k, idx in norm_dict.items():
                if "mms terminal id" in k or "merchant mms terminal id" in k:
                    col_map["mms_terminal_id"] = idx
                elif k == "terminal id" or k == "terminalid":
                    col_map["terminal_id"] = idx
                elif "partner ref id" in k or "partner ref" in k:
                    col_map["partner_ref_id"] = idx
                elif k == "vpa" or "vpa" in k:
                    col_map["vpa"] = idx
                elif "merchant id" in k:
                    col_map["merchant_id"] = idx
                elif "merchant name" in k:
                    col_map["merchant_name"] = idx
            break

    if header_row_idx is None or "mms_terminal_id" not in col_map:
        raise ValueError("Could not find required columns (MMS TERMINAL ID, Partner Ref ID / TERMINAL ID) in terminal file.")

    mappings_by_ref_id: Dict[str, Dict[str, str]] = {}
    mappings_by_vpa: Dict[str, str] = {}
    mappings_by_terminal_id: Dict[str, Dict[str, str]] = {}
    total_records = 0

    # Parse data rows
    for r_idx in range(header_row_idx + 1, len(sheet_rows)):
        row = sheet_rows[r_idx]
        mms_tid = clean_key(row[col_map["mms_terminal_id"]]) if "mms_terminal_id" in col_map and col_map["mms_terminal_id"] < len(row) else ""
        tid = clean_key(row[col_map["terminal_id"]]) if "terminal_id" in col_map and col_map["terminal_id"] < len(row) else ""
        pref_id = clean_key(row[col_map["partner_ref_id"]]) if "partner_ref_id" in col_map and col_map["partner_ref_id"] < len(row) else ""
        vpa = str(row[col_map["vpa"]]).strip() if "vpa" in col_map and col_map["vpa"] < len(row) and row[col_map["vpa"]] is not None else ""
        m_id = clean_key(row[col_map["merchant_id"]]) if "merchant_id" in col_map and col_map["merchant_id"] < len(row) else ""
        m_name = str(row[col_map["merchant_name"]]).strip() if "merchant_name" in col_map and col_map["merchant_name"] < len(row) and row[col_map["merchant_name"]] is not None else ""

        if not mms_tid and not tid:
            continue
        if mms_tid.lower() in ("mms terminal id", "terminal id", "mms_terminal_id", "merchant mms terminal id", "slno", "sr no"):
            continue

        effective_mms = mms_tid or tid
        effective_tid = tid or mms_tid

        record_payload = {
            "mms_terminal_id": effective_mms,
            "terminal_id": effective_tid,
            "partner_ref_id": pref_id,
            "vpa": vpa,
            "merchant_id": m_id,
            "merchant_name": m_name
        }
        total_records += 1

        # Index by Partner Ref ID (e.g. '4873', '4860', '4030', '4803', '138960')
        if pref_id and pref_id != "--":
            mappings_by_ref_id[pref_id] = record_payload

        # Index by VPA text or numbers within VPA (e.g. 'tn= 4803' -> '4803')
        if vpa and vpa != "--":
            clean_vpa = clean_key(vpa)
            if clean_vpa:
                mappings_by_vpa[clean_vpa] = effective_mms
            # Extract standalone numbers (e.g. 4873, 4803) from VPA string like 'tn= 4803'
            nums = re.findall(r'\b\d{3,8}\b', vpa)
            for n in nums:
                clean_n = clean_key(n)
                if clean_n and clean_n not in mappings_by_ref_id:
                    mappings_by_ref_id[clean_n] = record_payload
                if clean_n:
                    mappings_by_vpa[clean_n] = effective_mms

        # Index by Terminal ID
        if tid and tid != "--":
            mappings_by_terminal_id[tid] = record_payload

    result = {
        "updated_at": datetime.now().isoformat(),
        "source_filename": os.path.basename(file_path),
        "total_rows_scanned": total_records,
        "ref_id_count": len(mappings_by_ref_id),
        "terminal_id_count": len(mappings_by_terminal_id),
        "mappings_by_ref_id": mappings_by_ref_id,
        "mappings_by_vpa": mappings_by_vpa,
        "mappings_by_terminal_id": mappings_by_terminal_id
    }
    return result


def save_terminal_mappings(mappings_data: Dict[str, Any], merge: bool = True) -> str:
    """Saves parsed terminal mappings to disk, merging with existing mappings if merge=True."""
    global _CACHE
    os.makedirs(DATA_DIR, exist_ok=True)

    if merge and os.path.exists(MAPPINGS_FILE):
        try:
            with open(MAPPINGS_FILE, "r", encoding="utf-8") as f:
                existing = json.load(f)
                if isinstance(existing, dict):
                    existing_ref = existing.get("mappings_by_ref_id", {})
                    existing_ref.update(mappings_data.get("mappings_by_ref_id", {}))
                    mappings_data["mappings_by_ref_id"] = existing_ref

                    existing_vpa = existing.get("mappings_by_vpa", {})
                    existing_vpa.update(mappings_data.get("mappings_by_vpa", {}))
                    mappings_data["mappings_by_vpa"] = existing_vpa

                    existing_tid = existing.get("mappings_by_terminal_id", {})
                    existing_tid.update(mappings_data.get("mappings_by_terminal_id", {}))
                    mappings_data["mappings_by_terminal_id"] = existing_tid

                    mappings_data["ref_id_count"] = len(existing_ref)
                    mappings_data["terminal_id_count"] = len(existing_tid)
        except Exception as e:
            print(f"[TerminalMapper] Merge notice: {e}")

    with open(MAPPINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(mappings_data, f, indent=2)
    _CACHE = mappings_data
    return MAPPINGS_FILE


def load_terminal_mappings(force_reload: bool = False) -> Dict[str, Any]:
    """Loads terminal mappings from disk cache or memory."""
    global _CACHE
    if _CACHE is not None and not force_reload:
        return _CACHE

    if os.path.exists(MAPPINGS_FILE):
        try:
            with open(MAPPINGS_FILE, "r", encoding="utf-8") as f:
                _CACHE = json.load(f)
                return _CACHE
        except Exception:
            pass

    _CACHE = {
        "updated_at": None,
        "source_filename": None,
        "mappings_by_ref_id": {},
        "mappings_by_vpa": {},
        "mappings_by_terminal_id": {}
    }
    return _CACHE


def load_and_save_terminal_file(file_path: str, merge: bool = True) -> Dict[str, Any]:
    """Convenience function: parses an uploaded file and saves it immediately."""
    parsed = parse_terminal_report(file_path)
    save_terminal_mappings(parsed, merge=merge)
    return {
        "success": True,
        "source_filename": parsed.get("source_filename"),
        "total_mappings": len(parsed.get("mappings_by_ref_id", {})),
        "terminal_count": len(parsed.get("mappings_by_terminal_id", {})),
        "updated_at": parsed.get("updated_at")
    }


def extract_cf_middle_number(order_id: str) -> str:
    """
    Extracts the middle number from Cashfree Order IDs:
    e.g. '330595-4873-AXL706427102c4545c5a848c8186d8bd169axisupioffline' -> '4873'
    e.g. '330595-4860-AXIdbfc2143bcf04bd897057bc8f80e4eb7axisupioffline' -> '4860'
    e.g. '330595 4939 AXI3b51684e91814b13b3a5e2f986ab638bcfnsdlupioffline' -> '4939'
    """
    if not order_id:
        return ""
    s = str(order_id).strip()

    # Hyphen delimited: 330595-4873-...
    if "-" in s:
        parts = s.split("-")
        if len(parts) >= 2 and parts[1].strip():
            return clean_key(parts[1])

    # Space delimited: 330595 4939 ...
    if " " in s:
        parts = s.split()
        if len(parts) >= 2 and parts[1].strip():
            return clean_key(parts[1])

    # Regex fallback for \d+[- ](\d+)[- ]
    m = re.search(r'\d+[- ](\d+)[- ]', s)
    if m:
        return clean_key(m.group(1))

    return ""


def resolve_terminal_details(identifier: str, mappings: Optional[Dict[str, Any]] = None) -> Optional[Dict[str, Any]]:
    """
    Resolves an identifier (Cashfree Order ID, middle number, Partner Ref ID, Terminal ID, or pasted transaction row)
    to its full terminal mapping record: mms_terminal_id, terminal_id, partner_ref_id, merchant_name, vpa.
    """
    if not identifier:
        return None

    if mappings is None:
        mappings = load_terminal_mappings()

    by_ref = mappings.get("mappings_by_ref_id", {})
    by_vpa = mappings.get("mappings_by_vpa", {})
    by_tid = mappings.get("mappings_by_terminal_id", {})

    raw_str = str(identifier).strip()

    # If raw_str contains multi-column or multi-line text (e.g. pasted spreadsheet row), extract Order ID
    m_order = re.search(r'\b(\d{5,8}[- ]\d{3,8}[- ][A-Za-z0-9_]+)\b', raw_str)
    if m_order:
        raw_str = m_order.group(1).strip()

    mid = extract_cf_middle_number(raw_str)
    clean_id = clean_key(raw_str)

    candidates = [k for k in [mid, clean_id] if k]

    for cand in candidates:
        if cand in by_ref:
            rec = dict(by_ref[cand])
            rec["middle_number"] = mid or cand
            if not rec.get("mms_terminal_id") and rec.get("terminal_id"):
                rec["mms_terminal_id"] = rec["terminal_id"]
            return rec
        if cand in by_tid:
            rec = dict(by_tid[cand])
            rec["middle_number"] = mid or cand
            if not rec.get("mms_terminal_id") and rec.get("terminal_id"):
                rec["mms_terminal_id"] = rec["terminal_id"]
            return rec
        if cand in by_vpa:
            vpa_val = by_vpa[cand]
            return {
                "mms_terminal_id": vpa_val,
                "terminal_id": None,
                "partner_ref_id": cand,
                "middle_number": mid or cand,
                "vpa": cand,
                "merchant_name": ""
            }

    return None


def resolve_mms_terminal_id(identifier: str, mappings: Optional[Dict[str, Any]] = None) -> Optional[str]:
    """
    Resolves an identifier (Order ID, middle number, Partner Ref ID, or Terminal ID)
    to its corresponding MMS Terminal ID.
    """
    rec = resolve_terminal_details(identifier, mappings)
    if rec:
        return rec.get("mms_terminal_id") or rec.get("terminal_id")
    return None


def clear_terminal_mappings() -> bool:
    """Clears persisted mappings and resets in-memory cache."""
    global _CACHE
    if os.path.exists(MAPPINGS_FILE):
        try:
            os.remove(MAPPINGS_FILE)
        except Exception:
            pass
    _CACHE = {
        "updated_at": None,
        "source_filename": None,
        "mappings_by_ref_id": {},
        "mappings_by_vpa": {},
        "mappings_by_terminal_id": {}
    }
    return True
