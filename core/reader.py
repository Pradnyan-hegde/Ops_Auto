"""
File Reader for Excel (.xlsx, .xls) and CSV reports.
Preserves raw values, leading zeros, and distinguishes blanks from zeros.
"""
import os
import csv
import openpyxl
from dataclasses import dataclass
from typing import List, Dict, Any, Tuple, Optional
from .detector import ReportType, detect_report_type, validate_report_headers, MissingColumnsException


@dataclass
class RawReport:
    file_path: str
    report_type: ReportType
    headers: List[str]
    header_row_idx: int
    records: List[Dict[str, Any]]
    raw_matrix: List[List[Any]]  # all rows including header for preservation in raw tabs


def read_csv_report(file_path: str) -> RawReport:
    """Reads CSV file trying multiple encodings and extracts table data."""
    encodings = ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']
    lines = None
    used_encoding = 'utf-8'

    for enc in encodings:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                sample = f.read(4096)
                f.seek(0)
                lines = list(csv.reader(f))
                used_encoding = enc
                break
        except (UnicodeDecodeError, Exception):
            continue

    if lines is None:
        # Fallback with error ignore
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            lines = list(csv.reader(f))

    rep_type, header_idx, headers = detect_report_type(lines[:15])
    if rep_type == ReportType.UNKNOWN:
        # Try finding non-empty line
        for i, row in enumerate(lines[:10]):
            if any(cell.strip() for cell in row):
                header_idx = i
                headers = [c.strip() for c in row]
                break

    validate_report_headers(rep_type, headers, file_path)

    records: List[Dict[str, Any]] = []
    for r_idx in range(header_idx + 1, len(lines)):
        row = lines[r_idx]
        if not any(str(c).strip() for c in row if c is not None):
            continue
        first_c = str(row[0] or "").strip().lower()
        if first_c.startswith("report generated on") or first_c.startswith("total:"):
            continue
        rec = {}
        for c_idx, h in enumerate(headers):
            val = row[c_idx] if c_idx < len(row) else None
            rec[h] = val
        records.append(rec)

    return RawReport(
        file_path=file_path,
        report_type=rep_type,
        headers=headers,
        header_row_idx=header_idx,
        records=records,
        raw_matrix=lines
    )


def read_excel_report(file_path: str) -> RawReport:
    """Reads Excel (.xlsx / OOXML .xls / binary .xls) file."""
    import io

    with open(file_path, "rb") as f:
        file_bytes = f.read()

    magic = file_bytes[:8]
    wb = None
    raw_matrix: List[List[Any]] = []

    if magic.startswith(b"PK\x03\x04"):
        # OOXML / Zip-based Excel file (even if named .xls)
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=False, data_only=False)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            raw_matrix.append(list(row))
        wb.close()
    elif magic.startswith(b"\xd0\xcf\x11\xe0"):
        # Legacy BIFF8 binary Excel
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=file_bytes)
            sheet = book.sheet_by_index(0)
            raw_matrix = [sheet.row_values(r) for r in range(sheet.nrows)]
        except ImportError:
            raise RuntimeError(f"Cannot read legacy binary Excel file {file_path}. Please install xlrd.")
    else:
        # Try openpyxl directly
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            raw_matrix.append(list(row))
        wb.close()

    rep_type, header_idx, headers = detect_report_type(raw_matrix[:15])
    if rep_type == ReportType.UNKNOWN:
        # Check if other sheets have recognizable data
        wb2 = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
        for sname in wb2.sheetnames:
            ws_other = wb2[sname]
            cand_matrix = [list(r) for r in ws_other.iter_rows(values_only=True, max_row=15)]
            cand_type, cand_idx, cand_headers = detect_report_type(cand_matrix)
            if cand_type != ReportType.UNKNOWN:
                rep_type = cand_type
                header_idx = cand_idx
                headers = cand_headers
                # re-read entire sheet
                raw_matrix = [list(r) for r in ws_other.iter_rows(values_only=True)]
                break
        wb2.close()

    validate_report_headers(rep_type, headers, file_path)

    records = []
    for r_idx in range(header_idx + 1, len(raw_matrix)):
        row = raw_matrix[r_idx]
        if not any(c is not None and str(c).strip() != '' for c in row):
            continue
        first_c = str(row[0] or "").strip().lower()
        if first_c.startswith("report generated on") or first_c.startswith("total:"):
            continue
        rec = {}
        for c_idx, h in enumerate(headers):
            val = row[c_idx] if c_idx < len(row) else None
            rec[h] = val
        records.append(rec)

    return RawReport(
        file_path=file_path,
        report_type=rep_type,
        headers=headers,
        header_row_idx=header_idx,
        records=records,
        raw_matrix=raw_matrix
    )


def read_report(file_path: str) -> RawReport:
    """Entrypoint: Reads any supported report (.xlsx, .xls, .csv)."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Source file not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.csv':
        return read_csv_report(file_path)
    elif ext in ('.xlsx', '.xls', '.xlsm'):
        return read_excel_report(file_path)
    else:
        # Try CSV first, then Excel
        try:
            return read_csv_report(file_path)
        except Exception:
            return read_excel_report(file_path)


def read_all_reports_from_file(file_path: str) -> List[RawReport]:
    """Reads a file and returns all detected reports (extracts multiple sheets if present)."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Source file not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in ('.xlsx', '.xls', '.xlsm'):
        return [read_report(file_path)]

    results: List[RawReport] = []
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
        for sname in wb.sheetnames:
            ws = wb[sname]
            cand_matrix = [list(r) for r in ws.iter_rows(values_only=True)]
            if not cand_matrix:
                continue
            cand_type, cand_idx, cand_headers = detect_report_type(cand_matrix[:15])
            if cand_type != ReportType.UNKNOWN:
                try:
                    validate_report_headers(cand_type, cand_headers, f"{file_path} [{sname}]")
                    records = []
                    for r_idx in range(cand_idx + 1, len(cand_matrix)):
                        row = cand_matrix[r_idx]
                        if not any(c is not None and str(c).strip() != '' for c in row):
                            continue
                        first_c = str(row[0] or '').strip().lower()
                        if first_c.startswith('report generated on') or first_c.startswith('total:'):
                            continue
                        rec = {}
                        for c_idx, h in enumerate(cand_headers):
                            val = row[c_idx] if c_idx < len(row) else None
                            rec[h] = val
                        records.append(rec)
                    results.append(RawReport(
                        file_path=file_path,
                        report_type=cand_type,
                        headers=cand_headers,
                        header_row_idx=cand_idx,
                        records=records,
                        raw_matrix=cand_matrix
                    ))
                except Exception:
                    pass
        wb.close()
    except Exception:
        pass

    if results:
        return results

    return [read_report(file_path)]

