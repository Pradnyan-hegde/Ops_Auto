"""
SMMS Sync File Builder.
Generates an Excel file for transactions present in CMS but missing in SMMS
so they can be imported and synchronized in SMMS.
Matching format shown in sample 'Sync Transaction2 (1).xlsx':
- Column A: SwinkPay Transaction ID
- No borders, clean styling.
"""
import os
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, Alignment


def build_smms_sync_workbook(
    cms_not_in_smms: List[Dict[str, Any]],
    output_path: str
) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    font_header = Font(name="Calibri", size=11, bold=True, color="000000")
    font_data = Font(name="Calibri", size=11, bold=False, color="000000")
    align_left = Alignment(horizontal="left", vertical="center")

    # Header in A1
    c1 = ws.cell(1, 1, "SwinkPay Transaction ID")
    c1.font = font_header
    c1.alignment = align_left

    # Rows 2+
    for row_idx, r in enumerate(cms_not_in_smms, start=2):
        sp_id = str(r.get("SwinkPay Txn ID") or r.get("SwinkPay Transaction ID") or r.get("Matched Transaction ID") or "").strip()
        cell = ws.cell(row_idx, 1, sp_id)
        cell.font = font_data
        cell.alignment = align_left
        cell.number_format = "@"

    ws.column_dimensions["A"].width = 30

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path