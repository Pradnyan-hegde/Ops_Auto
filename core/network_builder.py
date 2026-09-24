"""
Change Network Builder for SwinkPay Reconciliation.
Detects discrepancies where CMS recorded an incorrect network (e.g. 'RUPAY')
while Cashfree report captured the true payment mode (e.g. 'UPI_CREDIT_CARD_OFFLINE_STATIC').
Builds 'Change Network.xlsx' workbook matching sample format:
- A1: SwinkPay Transaction ID
- B1: Old Network
- C1: New Network
- Borderless styling, Calibri 11pt.
"""
import os
import re
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, Alignment

from .normalizer import clean_key


def normalize_network_name(val: Any) -> str:
    """Normalizes network / payment mode string for accurate comparison."""
    if val is None:
        return ""
    s = str(val).strip().lower()
    # Normalize whitespace, underscores, hyphens
    s = re.sub(r"[\s_\-]+", "_", s)
    return s


def are_networks_equivalent(cms_net: str, cf_mode: str) -> bool:
    """
    Determines whether a CMS Network string and Cashfree Payment Mode represent
    the same payment method (i.e. no network change needed).

    Equivalencies:
    - Standard UPI: 'upi', 'upi_offline_static', 'upi_qr', 'static_qr', 'upi_offline'
    - Credit Card on UPI: 'upi_credit_card_offline_static', 'upi_cc', 'upi_credit_card'
    - PPI / Wallet on UPI: 'upi_ppi_offline_static', 'upi_ppi', 'upi_wallet'

    Discrepancies (Returns False):
    - CMS 'RUPAY' vs CF 'UPI_CREDIT_CARD_OFFLINE_STATIC' (Card misclassification)
    - CMS 'WA' vs CF 'UPI_PPI_OFFLINE_STATIC' (Wallet misclassification)
    """
    n_cms = normalize_network_name(cms_net)
    n_cf = normalize_network_name(cf_mode)

    if not n_cms or not n_cf:
        return True

    if n_cms == n_cf:
        return True

    std_upi = {"upi", "upi_offline_static", "upi_qr", "static_qr", "upi_offline"}
    if n_cms in std_upi and n_cf in std_upi:
        return True

    cc_upi = {"upi_credit_card_offline_static", "upi_cc", "upi_credit_card"}
    if n_cms in cc_upi and n_cf in cc_upi:
        return True

    ppi_upi = {"upi_ppi_offline_static", "upi_ppi", "upi_wallet"}
    if n_cms in ppi_upi and n_cf in ppi_upi:
        return True

    return False


def detect_network_changes(
    matched_cf_records: Optional[List[Dict[str, Any]]] = None,
    cms_records: Optional[List[Dict[str, Any]]] = None,
    cf_records: Optional[List[Dict[str, Any]]] = None,
    cms_not_in_smms: Optional[List[Dict[str, Any]]] = None
) -> List[Dict[str, str]]:
    """
    Identifies transactions where CMS Network differs from Cashfree Payment Mode.
    Returns a list of dicts:
    [
        {
            "SwinkPay Transaction ID": "SPRpeCojRXmGOc05gLFA",
            "Old Network": "RUPAY",
            "New Network": "UPI_CREDIT_CARD_OFFLINE_STATIC"
        },
        ...
    ]
    """
    # Build lookup map for Cashfree records by Bank Reference No or ID
    cf_by_ref: Dict[str, Dict[str, Any]] = {}
    if cf_records:
        for r in cf_records:
            ref = clean_key(
                r.get("Bank Reference No.") or r.get("Bank Reference No") or r.get("Reference Id") or r.get("Order Id")
            ).lstrip("0")
            if ref:
                cf_by_ref[ref] = r

    # Determine candidate CMS records
    candidate_recs: List[Dict[str, Any]] = []
    if cms_records:
        candidate_recs.extend(cms_records)
    if matched_cf_records:
        candidate_recs.extend(matched_cf_records)
    if cms_not_in_smms:
        candidate_recs.extend(cms_not_in_smms)

    changes: List[Dict[str, str]] = []
    seen_sp_ids = set()

    for r in candidate_recs:
        sp_id = str(
            r.get("SwinkPay Txn ID") or r.get("SwinkPay Transaction ID") or r.get("Matched Transaction ID") or ""
        ).strip()
        if not sp_id or sp_id in seen_sp_ids:
            continue

        cms_network = str(r.get("Network") or "").strip()

        # Try to find Cashfree payment mode
        cf_mode = str(r.get("Cashfree Payment Mode") or r.get("Payment Mode") or "").strip()
        if not cf_mode and cf_by_ref:
            ref_key = clean_key(r.get("RRN/UTR") or r.get("Match Key") or r.get("RRN")).lstrip("0")
            cf_match = cf_by_ref.get(ref_key)
            if cf_match:
                cf_mode = str(cf_match.get("Payment Mode") or "").strip()

        if not cf_mode:
            continue

        # If networks are not equivalent, record change
        if not are_networks_equivalent(cms_network, cf_mode):
            seen_sp_ids.add(sp_id)
            changes.append({
                "SwinkPay Transaction ID": sp_id,
                "Old Network": cms_network,
                "New Network": cf_mode
            })

    return changes


def build_change_network_workbook(
    network_changes: List[Dict[str, str]],
    output_path: str
) -> str:
    """
    Builds an Excel workbook (.xlsx) with:
    - Sheet1
    - Row 1: SwinkPay Transaction ID, Old Network, New Network
    - Rows 2+: Values
    - Calibri 11pt, left-aligned, text formatted, borderless.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    font_header = Font(name="Calibri", size=11, bold=True, color="000000")
    font_data = Font(name="Calibri", size=11, bold=False, color="000000")
    align_left = Alignment(horizontal="left", vertical="center")

    # Header Row
    headers = ["SwinkPay Transaction ID", "Old Network", "New Network"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, h)
        cell.font = font_header
        cell.alignment = align_left

    # Data Rows
    for row_idx, item in enumerate(network_changes, start=2):
        c_sp = ws.cell(row_idx, 1, item.get("SwinkPay Transaction ID", ""))
        c_sp.font = font_data
        c_sp.alignment = align_left
        c_sp.number_format = "@"

        c_old = ws.cell(row_idx, 2, item.get("Old Network", ""))
        c_old.font = font_data
        c_old.alignment = align_left
        c_old.number_format = "@"

        c_new = ws.cell(row_idx, 3, item.get("New Network", ""))
        c_new.font = font_data
        c_new.alignment = align_left
        c_new.number_format = "@"

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 35

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
