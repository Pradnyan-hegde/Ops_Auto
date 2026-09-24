"""
Automated Verification Tests for Ops_Auto.
Verifies header detection, missing columns handling, matching logic, and output fidelity.
"""
import os
import sys
import unittest
import openpyxl

# Add parent directory to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from core.detector import ReportType, detect_report_type, validate_report_headers, MissingColumnsException
from core.normalizer import clean_key, clean_amount, is_amount_equal, normalize_status
from core.reader import read_report
from core.matcher import ReconciliationEngine
from core.aggregator import Aggregator
from core.excel_builder import ExcelReportBuilder


class TestNormalizer(unittest.TestCase):
    def test_clean_key_leading_zeros(self):
        # Must never lose leading zero
        self.assertEqual(clean_key("'090984560431"), "090984560431")
        self.assertEqual(clean_key("090984560431"), "090984560431")
        self.assertEqual(clean_key("000123"), "000123")

    def test_clean_key_floats(self):
        # Float strings ending in .0
        self.assertEqual(clean_key("472695706519.0"), "472695706519")
        self.assertEqual(clean_key(472695706519.0), "472695706519")

    def test_clean_amount_distinguishes_zero_and_none(self):
        # 0 must be 0.0, None / empty must be None
        self.assertIsNone(clean_amount(None))
        self.assertIsNone(clean_amount(""))
        self.assertIsNone(clean_amount("--"))
        self.assertEqual(clean_amount(0), 0.0)
        self.assertEqual(clean_amount("0.00"), 0.0)
        self.assertEqual(clean_amount("Rs. 1,250.50"), 1250.50)

    def test_is_amount_equal(self):
        self.assertTrue(is_amount_equal(100.0, 100.005, tolerance=0.01))
        self.assertFalse(is_amount_equal(100.0, 100.05, tolerance=0.01))
        self.assertFalse(is_amount_equal(None, 0.0))
        self.assertTrue(is_amount_equal(None, None))

    def test_normalize_status_codes(self):
        self.assertEqual(normalize_status(1), "Success")
        self.assertEqual(normalize_status("1"), "Success")
        self.assertEqual(normalize_status(2), "Failed")
        self.assertEqual(normalize_status("2"), "Failed")
        self.assertEqual(normalize_status(3), "Status 3")
        self.assertEqual(normalize_status("3"), "Status 3")
        self.assertEqual(normalize_status(4), "Status 4")
        self.assertEqual(normalize_status("4"), "Status 4")
        self.assertEqual(normalize_status(5), "Status 5")
        self.assertEqual(normalize_status("5"), "Status 5")
        self.assertEqual(normalize_status(6), "Status 6")
        self.assertEqual(normalize_status("6"), "Status 6")


class TestHeaderValidation(unittest.TestCase):
    def test_missing_column_raises(self):
        # Missing 'Bank Reference No.' in Cashfree
        invalid_cf_headers = ["Amount", "Settlement Amount", "Payment Mode", "Transaction Status"]
        with self.assertRaises(MissingColumnsException) as ctx:
            validate_report_headers(ReportType.CASHFREE, invalid_cf_headers)
        self.assertIn("Bank Reference No.", str(ctx.exception))


class TestEndToEndReconciliation(unittest.TestCase):
    def test_real_data_reconciliation(self):
        sample_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "sample_inputs"))
        cms_file = os.path.join(sample_dir, "CMS_Report_2026-09-10.xlsx")
        smms_file = os.path.join(sample_dir, "SMMS_Report_2026-09-10.xlsx")
        cf_file = os.path.join(sample_dir, "Cashfree_Report_2026-09-10.xlsx")
        eb_file = os.path.join(sample_dir, "Easebuzz_Report_2026-09-10.csv")
        air_file = os.path.join(sample_dir, "Airtel_Report_2026-09-10.csv")
        settle_file = os.path.join(sample_dir, "Airtel_Settlement_2026-09-10.csv")

        for p in [cms_file, smms_file, cf_file, eb_file, air_file]:
            if not os.path.exists(p):
                self.skipTest(f"Sample file not found: {p}")

        cms_rep = read_report(cms_file)
        smms_rep = read_report(smms_file)
        cf_rep = read_report(cf_file)
        eb_rep = read_report(eb_file)
        air_rep = read_report(air_file)
        settle_rep = read_report(settle_file) if os.path.exists(settle_file) else None

        engine = ReconciliationEngine(tolerance=0.01)
        engine.set_reports(
            cms=cms_rep,
            smms=smms_rep,
            cf=cf_rep,
            eb=eb_rep,
            airtel=air_rep,
            settle=settle_rep
        )
        engine.run()

        aggregator = Aggregator(engine)
        aggregator.compute()

        # Verify ground truth counts matching historical reference
        self.assertEqual(aggregator.grand_total["success_count"], 6515)
        self.assertEqual(aggregator.grand_total["gross_amount"], 434115.00)
        self.assertEqual(aggregator.grand_total["net_amount"], 428404.13)
        self.assertEqual(len(engine.failed_or_reversed), 93)

        # Partner counts
        cf_sub = aggregator.partner_subtotals.get("CashFree")
        self.assertIsNotNone(cf_sub)
        self.assertEqual(cf_sub["success_count"], 2691)
        self.assertEqual(cf_sub["matched_count"], 2691)

        eb_sub = aggregator.partner_subtotals.get("EaseBuzz")
        self.assertIsNotNone(eb_sub)
        self.assertEqual(eb_sub["success_count"], 2873)
        self.assertEqual(eb_sub["matched_count"], 2873)

        air_sub = aggregator.partner_subtotals.get("Airtel Bank")
        self.assertIsNotNone(air_sub)
        self.assertEqual(air_sub["success_count"], 951)
        self.assertEqual(air_sub["matched_count"], 951)

        # Difference-to-zero control checks
        for chk in aggregator.control_checks:
            self.assertEqual(chk["status"], "PASS", f"Failed control check: {chk['check_name']} diff={chk['difference']}")

        # Build and verify Excel workbook
        out_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "test_reconciliation_output.xlsx"))
        builder = ExcelReportBuilder(engine, aggregator)
        builder.build_and_save(out_path)
        self.assertTrue(os.path.exists(out_path))

        # Check sheets in workbook
        wb = openpyxl.load_workbook(out_path, read_only=True)
        expected_sheets = [
            "CMS", "SMMS", "Easebuzz", "Cashfree", "Airtel", "cms_failure_txn",
            "CMS_SMMS_Matched", "CMS_CF_Matched", "CMS_EB_Matched", "CMS_Air_Matched",
            "Summary"
        ]
        for s in expected_sheets:
            self.assertIn(s, wb.sheetnames)
        self.assertEqual(wb.sheetnames[-1], "Summary", "Summary must be the last sheet")
        wb.close()

        # Clean up test output
        try:
            os.remove(out_path)
        except OSError:
            pass

    def test_cms_status_3_4_5_6_sheets(self):
        from core.reader import RawReport
        cms_headers = ["Sl. No.", "SwinkPay Txn ID", "RRN/UTR", "Transaction Amount", "Transaction Status", "Network"]
        cms_records = [
            {"Sl. No.": 1, "SwinkPay Txn ID": "SP1", "RRN/UTR": "1001", "Transaction Amount": 100.0, "Transaction Status": 1, "Network": "UPI"},
            {"Sl. No.": 2, "SwinkPay Txn ID": "SP2", "RRN/UTR": "1002", "Transaction Amount": 200.0, "Transaction Status": 2, "Network": "UPI"},
            {"Sl. No.": 3, "SwinkPay Txn ID": "SP3", "RRN/UTR": "1003", "Transaction Amount": 300.0, "Transaction Status": 3, "Network": "UPI"},
            {"Sl. No.": 4, "SwinkPay Txn ID": "SP4", "RRN/UTR": "1004", "Transaction Amount": 400.0, "Transaction Status": 4, "Network": "UPI"},
            {"Sl. No.": 5, "SwinkPay Txn ID": "SP5", "RRN/UTR": "1005", "Transaction Amount": 500.0, "Transaction Status": 5, "Network": "UPI"},
            {"Sl. No.": 6, "SwinkPay Txn ID": "SP6", "RRN/UTR": "1006", "Transaction Amount": 600.0, "Transaction Status": 6, "Network": "UPI"},
        ]
        raw_matrix = [cms_headers] + [[r[h] for h in cms_headers] for r in cms_records]
        cms_rep = RawReport(
            file_path="cms.xlsx",
            report_type=ReportType.CMS,
            headers=cms_headers,
            header_row_idx=0,
            records=cms_records,
            raw_matrix=raw_matrix
        )

        smms_records = [
            {"Sl. No.": 1, "SwinkPay Txn ID": "SP1", "RRN/UTR": "1001", "Transaction Amount": 100.0, "Transaction Status": "Success", "Network": "UPI", "PG/Bank": "CashFree"}
        ]
        smms_rep = RawReport(
            file_path="smms.xlsx",
            report_type=ReportType.SMMS,
            headers=cms_headers + ["PG/Bank"],
            header_row_idx=0,
            records=smms_records,
            raw_matrix=[cms_headers + ["PG/Bank"]]
        )

        engine = ReconciliationEngine()
        engine.set_reports(cms=cms_rep, smms=smms_rep)
        engine.run()

        aggregator = Aggregator(engine)
        aggregator.compute()

        out_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "test_status_sheets.xlsx"))
        builder = ExcelReportBuilder(engine, aggregator)
        builder.build_and_save(out_path)

        wb = openpyxl.load_workbook(out_path, read_only=True)
        self.assertIn("cms_failure_txn", wb.sheetnames)
        self.assertIn("cms_status_3_txn", wb.sheetnames)
        self.assertIn("cms_status_4_txn", wb.sheetnames)
        self.assertIn("cms_status_5_txn", wb.sheetnames)
        self.assertIn("cms_status_6_txn", wb.sheetnames)
        self.assertEqual(wb.sheetnames[-1], "Summary")
        wb.close()
        os.remove(out_path)


class TestXCDDownloads(unittest.TestCase):
    def setUp(self):
        self.test_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "scratch_test_xcd"))
        os.makedirs(self.test_dir, exist_ok=True)

    def tearDown(self):
        import shutil
        if os.path.exists(self.test_dir):
            shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_xcd_workbook_format_and_integrity(self):
        from core.xcd_builder import build_xcd_workbook
        from core.xcd_validator import validate_xcd_workbook_file

        records = [
            {"SwinkPay Txn ID": "SP_TXN_001", "Transaction Amount": 150.00},
            {"SwinkPay Txn ID": "SP_TXN_002", "Transaction Amount": 250.50},
            {"SwinkPay Txn ID": "SP_TXN_003", "Transaction Amount": 75.00},
        ]
        out_file = os.path.join(self.test_dir, "test_xcd.xlsx")
        build_xcd_workbook(records, "11-09-2026", out_file)

        self.assertTrue(os.path.exists(out_file))

        # Check physical workbook structure
        wb = openpyxl.load_workbook(out_file, data_only=True)
        ws = wb.active
        self.assertEqual(ws.cell(1, 1).value, "SALE TRANSACTIONS")
        headers = [ws.cell(2, c).value for c in range(1, 5)]
        self.assertEqual(headers, ["SL NO", "SwinkPay Transaction ID", "Amount", "Settlement Date"])
        self.assertEqual(ws.cell(3, 1).value, 1)
        self.assertEqual(ws.cell(3, 2).value, "SP_TXN_001")
        self.assertEqual(ws.cell(3, 3).value, 150.0)
        self.assertEqual(ws.cell(3, 4).value, "11-09-2026")
        self.assertEqual(ws.cell(5, 1).value, 3)
        self.assertEqual(ws.cell(5, 2).value, "SP_TXN_003")
        wb.close()

        # Run independent validator
        val_res = validate_xcd_workbook_file(out_file, records)
        self.assertTrue(val_res["valid"])
        self.assertEqual(val_res["row_count"], 3)
        self.assertEqual(val_res["total_amount"], 475.50)

    def test_all_clear_gate_pass_and_block_conditions(self):
        from core.xcd_validator import validate_all_clear

        class MockAggregator:
            control_checks = [{"check_name": "Check 1", "status": "PASS"}]

        class MockEngine:
            tolerance = 0.01
            cms_not_in_smms = []
            smms_not_in_cms = []
            unmatched_cf = []
            unmatched_eb = []
            unmatched_airtel = []
            duplicates = []
            exceptions = []
            cms_cf_matched = [{"CMS Amount": 100.0, "Partner Amount": 100.0, "CMS Status": "Success", "Partner Status": "Success"}]
            cms_eb_matched = [{"CMS Amount": 200.0, "Partner Amount": 200.0, "CMS Status": "Success", "Partner Status": "Success"}]
            cms_air_matched = [{"CMS Amount": 300.0, "Partner Amount": 300.0, "CMS Status": "Success", "Partner Status": "Success"}]

        engine = MockEngine()
        agg = MockAggregator()

        # 1. Clean run -> PASS
        res = validate_all_clear(engine, agg, settlement_date="11-09-2026")
        self.assertTrue(res.is_all_clear)
        self.assertEqual(res.summary_message, "Reconciliation complete. All three XCD settlement files are ready to download.")

        # 2. Unmatched Airtel -> BLOCKED
        engine.unmatched_airtel = [{"Match Key": "12345"}]
        res_blocked = validate_all_clear(engine, agg, settlement_date="11-09-2026")
        self.assertFalse(res_blocked.is_all_clear)
        self.assertTrue(any("unmatched Airtel" in r for r in res_blocked.blocking_reasons))
        engine.unmatched_airtel = []

        # 3. Amount mismatch -> BLOCKED
        engine.cms_cf_matched[0]["Partner Amount"] = 99.0
        res_amt_mismatch = validate_all_clear(engine, agg, settlement_date="11-09-2026")
        self.assertFalse(res_amt_mismatch.is_all_clear)
        self.assertTrue(any("amount mismatch" in r for r in res_amt_mismatch.blocking_reasons))
        engine.cms_cf_matched[0]["Partner Amount"] = 100.0

        # 4. Missing settlement date -> BLOCKED
        res_no_date = validate_all_clear(engine, agg, settlement_date="")
        self.assertFalse(res_no_date.is_all_clear)
        self.assertTrue(any("Settlement date is missing" in r for r in res_no_date.blocking_reasons))


if __name__ == "__main__":
    unittest.main()

