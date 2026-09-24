"""
Unit and integration tests for SMMS Sync File generator and
Payment Gateway Missing Transaction Postman Payload builder.
"""
import os
import sys
import json
import unittest
import openpyxl

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from core.sync_builder import build_smms_sync_workbook
from core.pg_payload_builder import (
    format_payload_datetime,
    format_payload_amount,
    build_cms_prefix_terminal_map,
    build_missing_pg_payloads,
)
from dashboard.server import download_sync_file, SESSIONS_DIR


class DummyReport:
    def __init__(self, records):
        self.records = records


class DummyEngine:
    def __init__(self, cms_report=None, unmatched_cf=None, unmatched_eb=None, unmatched_airtel=None, cms_not_in_smms=None):
        self.cms_report = cms_report
        self.unmatched_cf = unmatched_cf or []
        self.unmatched_eb = unmatched_eb or []
        self.unmatched_airtel = unmatched_airtel or []
        self.cms_not_in_smms = cms_not_in_smms or []


class TestSyncAndPayloads(unittest.TestCase):
    def test_smms_sync_file_builder(self):
        records = [
            {"SwinkPay Txn ID": "SPUpBxEGwo4EHR4xicaB"},
            {"SwinkPay Txn ID": "SPUp9999999999999999"},
        ]
        test_path = os.path.join("scratch", "test_sync_output.xlsx")
        build_smms_sync_workbook(records, test_path)

        self.assertTrue(os.path.exists(test_path))
        wb = openpyxl.load_workbook(test_path)
        ws = wb.active

        # Header check
        self.assertEqual(ws.cell(1, 1).value, "SwinkPay Transaction ID")
        # Row data check
        self.assertEqual(ws.cell(2, 1).value, "SPUpBxEGwo4EHR4xicaB")
        self.assertEqual(ws.cell(3, 1).value, "SPUp9999999999999999")
        self.assertEqual(ws.max_row, 3)

        # Check no borders
        c1_b = ws.cell(1, 1).border
        self.assertFalse(bool(c1_b and c1_b.left and c1_b.left.style))

    def test_payload_formatters(self):
        # Datetime
        self.assertEqual(format_payload_datetime("2026-09-17 15:59:29"), "2026-09-17 15:59:29")
        self.assertEqual(format_payload_datetime("17-09-2026 15:59:29"), "2026-09-17 15:59:29")
        self.assertEqual(format_payload_datetime("17/09/26 15:59:29"), "2026-09-17 15:59:29")

        # Amount
        self.assertEqual(format_payload_amount(20), "20.00")
        self.assertEqual(format_payload_amount("20.5"), "20.50")
        self.assertEqual(format_payload_amount("1,250.75"), "1250.75")

    def test_cashfree_terminal_resolution(self):
        cms_records = [
            {
                "Merchant MMS Terminal ID": "X1JYYO",
                "Auth Response": '{"data":{"order":{"order_id":"330595-4873-PTMA8FA73937F8D4E79ACBF14A4BB7F99BAaxisupioffline"}}}',
            },
            {
                "Merchant MMS Terminal ID": "XKD8M3",
                "Auth Response": '{"data":{"order":{"order_id":"330595-5205-OMS99395edc07d248df8b6a97f60e94fd96axisupioffline"}}}',
            }
        ]
        cms_report = DummyReport(cms_records)
        prefix_map = build_cms_prefix_terminal_map(cms_report)
        self.assertEqual(prefix_map.get("330595-4873"), "X1JYYO")
        self.assertEqual(prefix_map.get("330595-5205"), "XKD8M3")

        # Unmatched CF record
        unmatched_cf = [
            {
                "Order Id": "330595-4873-SBI1d3b001fc32d494587fface9454b383eaxisupioffline",
                "Bank Reference No.": "129761680150",
                "Amount": 20.0,
                "Transaction Time": "2026-09-17 15:59:29"
            }
        ]
        engine = DummyEngine(cms_report=cms_report, unmatched_cf=unmatched_cf)
        payloads = build_missing_pg_payloads(engine)
        self.assertEqual(len(payloads), 1)

        p = payloads[0]
        self.assertEqual(p["gateway"], "Cashfree")
        self.assertEqual(p["terminal_id"], "X1JYYO")
        self.assertEqual(p["utr"], "129761680150")
        self.assertEqual(p["amount"], "20.00")
        self.assertEqual(p["date_and_time"], "2026-09-17 15:59:29")

        # Verify exact JSON keys match Postman schema
        parsed_json = json.loads(p["payload_json"])
        self.assertEqual(parsed_json["amount"], "20.00")
        self.assertEqual(parsed_json["terminalID"], "X1JYYO")
        self.assertEqual(parsed_json["utr"], "129761680150")
        self.assertEqual(parsed_json["dateAndTime"], "2026-09-17 15:59:29")

    def test_easebuzz_terminal_resolution(self):
        unmatched_eb = [
            {
                "ID": "EB_TXN_01",
                "UTR": "999888777666",
                "Amount": 150.0,
                "Transaction Date": "10-09-2026 12:30:00",
                "Virtual Account Label": "XAFNSG"
            }
        ]
        engine = DummyEngine(unmatched_eb=unmatched_eb)
        payloads = build_missing_pg_payloads(engine)
        self.assertEqual(len(payloads), 1)
        p = payloads[0]
        self.assertEqual(p["gateway"], "Easebuzz")
        self.assertEqual(p["terminal_id"], "XAFNSG")
        self.assertEqual(p["utr"], "999888777666")
        self.assertEqual(p["amount"], "150.00")
        self.assertEqual(p["date_and_time"], "2026-09-10 12:30:00")

    def test_airtel_terminal_resolution(self):
        unmatched_airtel = [
            {
                "Transaction Id": "AIR_TXN_01",
                "PARTNER_TXN_ID": "RRN11223344",
                "Original Input Amt": 75.5,
                "Date and Time": "10/09/26 14:20:00",
                "Transaction To": "spf-2045-xh2vuy@mairtel"
            }
        ]
        engine = DummyEngine(unmatched_airtel=unmatched_airtel)
        payloads = build_missing_pg_payloads(engine)
        self.assertEqual(len(payloads), 1)
        p = payloads[0]
        self.assertEqual(p["gateway"], "Airtel")
        self.assertEqual(p["terminal_id"], "XH2VUY")
        self.assertEqual(p["utr"], "RRN11223344")
        self.assertEqual(p["amount"], "75.50")
        self.assertEqual(p["date_and_time"], "2026-09-10 14:20:00")

    def test_download_sync_file_endpoint(self):
        session_id = "test_sync_session_789"
        session_dir = os.path.join(SESSIONS_DIR, session_id)
        os.makedirs(session_dir, exist_ok=True)
        try:
            sync_file = os.path.join(session_dir, "Sync_Transactions_2026-09-17.xlsx")
            build_smms_sync_workbook([{"SwinkPay Txn ID": "SPTEST1"}], sync_file)

            resp = download_sync_file(session_id)
            self.assertEqual(resp.filename, "Sync_Transactions_2026-09-17.xlsx")
            self.assertEqual(resp.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        finally:
            import shutil
            shutil.rmtree(session_dir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()