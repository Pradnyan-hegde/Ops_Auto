"""
Unit & Integration Tests for Multi-Day Weekend Reconciliation,
Date Checkbox Filtering, Reconciliation Workbook Parsing,
and Change Network File Automation.
"""
import os
import json
import shutil
import unittest
import openpyxl
from starlette.datastructures import UploadFile

from core.network_builder import (
    detect_network_changes,
    build_change_network_workbook,
    normalize_network_name,
    are_networks_equivalent
)
from core.xcd_builder import (
    get_distinct_dates_with_counts,
    extract_record_date,
    build_xcd_workbook,
    format_settlement_date_display,
)
from core.recon_parser import parse_reconciliation_workbook
from dashboard.server import (
    upload_recon_file,
    download_xcd,
    download_network_file,
    SESSIONS_DIR,
)


class TestMultiDayAndNetwork(unittest.TestCase):

    def setUp(self):
        self.test_dir = os.path.join(os.path.dirname(__file__), "test_scratch")
        os.makedirs(self.test_dir, exist_ok=True)

    def test_normalize_network_name(self):
        self.assertEqual(normalize_network_name("UPI_OFFLINE_STATIC"), "upi_offline_static")
        self.assertEqual(normalize_network_name("upi offline static"), "upi_offline_static")
        self.assertEqual(normalize_network_name("UPI-CREDIT-CARD"), "upi_credit_card")
        self.assertEqual(normalize_network_name(None), "")

    def test_are_networks_equivalent(self):
        # Standard UPI variations are equivalent (no network change)
        self.assertTrue(are_networks_equivalent("UPI", "UPI_OFFLINE_STATIC"))
        self.assertTrue(are_networks_equivalent("upi_offline_static", "UPI_OFFLINE_STATIC"))
        self.assertTrue(are_networks_equivalent("UPI", "upi"))

        # Credit card UPI variations are equivalent
        self.assertTrue(are_networks_equivalent("upi_credit_card_offline_static", "UPI_CREDIT_CARD_OFFLINE_STATIC"))
        self.assertTrue(are_networks_equivalent("upi_cc", "UPI_CREDIT_CARD_OFFLINE_STATIC"))

        # PPI / Wallet UPI variations are equivalent
        self.assertTrue(are_networks_equivalent("upi_ppi_offline_static", "UPI_PPI_OFFLINE_STATIC"))
        self.assertTrue(are_networks_equivalent("upi_ppi", "UPI_PPI_OFFLINE_STATIC"))

        # Genuine mismatches must return False
        self.assertFalse(are_networks_equivalent("RUPAY", "UPI_CREDIT_CARD_OFFLINE_STATIC"))
        self.assertFalse(are_networks_equivalent("WA", "UPI_PPI_OFFLINE_STATIC"))

    def test_detect_network_changes(self):
        # CMS records with standard UPI (should NOT change), Rupay CC (should change), Wallet (should change)
        matched_records = [
            {
                "SwinkPay Txn ID": "SP_RUPAY_1",
                "Network": "RUPAY",
                "Cashfree Payment Mode": "UPI_CREDIT_CARD_OFFLINE_STATIC"
            },
            {
                "SwinkPay Txn ID": "SP_WA_1",
                "Network": "WA",
                "Cashfree Payment Mode": "UPI_PPI_OFFLINE_STATIC"
            },
            {
                "SwinkPay Txn ID": "SP_UPI_NORMAL_1",
                "Network": "UPI",
                "Cashfree Payment Mode": "UPI_OFFLINE_STATIC"
            },
            {
                "SwinkPay Txn ID": "SP_UPI_NORMAL_2",
                "Network": "upi_offline_static",
                "Cashfree Payment Mode": "UPI_OFFLINE_STATIC"
            },
            {
                "SwinkPay Txn ID": "SP_CC_NORMAL_1",
                "Network": "upi_credit_card_offline_static",
                "Cashfree Payment Mode": "UPI_CREDIT_CARD_OFFLINE_STATIC"
            }
        ]

        changes = detect_network_changes(matched_records)
        self.assertEqual(len(changes), 2)
        self.assertEqual(changes[0]["SwinkPay Transaction ID"], "SP_RUPAY_1")
        self.assertEqual(changes[0]["Old Network"], "RUPAY")
        self.assertEqual(changes[0]["New Network"], "UPI_CREDIT_CARD_OFFLINE_STATIC")
        self.assertEqual(changes[1]["SwinkPay Transaction ID"], "SP_WA_1")
        self.assertEqual(changes[1]["Old Network"], "WA")
        self.assertEqual(changes[1]["New Network"], "UPI_PPI_OFFLINE_STATIC")

    def test_build_change_network_workbook(self):
        changes = [
            {
                "SwinkPay Transaction ID": "SPRpeCojRXmGOc05gLFA",
                "Old Network": "RUPAY",
                "New Network": "UPI_CREDIT_CARD_OFFLINE_STATIC"
            }
        ]
        out_path = os.path.join(self.test_dir, "test_change_network.xlsx")
        build_change_network_workbook(changes, out_path)

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        self.assertEqual(ws.title, "Sheet1")
        # Row 1 headers
        self.assertEqual(ws.cell(1, 1).value, "SwinkPay Transaction ID")
        self.assertEqual(ws.cell(1, 2).value, "Old Network")
        self.assertEqual(ws.cell(1, 3).value, "New Network")
        # Row 2 data
        self.assertEqual(ws.cell(2, 1).value, "SPRpeCojRXmGOc05gLFA")
        self.assertEqual(ws.cell(2, 2).value, "RUPAY")
        self.assertEqual(ws.cell(2, 3).value, "UPI_CREDIT_CARD_OFFLINE_STATIC")

        # Verify borderless styling
        for r in ws.iter_rows():
            for c in r:
                self.assertIsNone(c.border.left.style)
                self.assertIsNone(c.border.right.style)
                self.assertIsNone(c.border.top.style)
                self.assertIsNone(c.border.bottom.style)

    def test_distinct_dates_with_counts(self):
        recs_cf = [
            {"Transaction Date & Time": "2026-09-18 10:00:00"},
            {"Transaction Date & Time": "2026-09-18 11:00:00"},
            {"Transaction Date & Time": "2026-09-19 12:00:00"},
        ]
        recs_eb = [
            {"Transaction Date & Time": "2026-09-18 10:30:00"},
            {"Transaction Date & Time": "2026-09-20 15:00:00"},
        ]
        recs_air = [
            {"Transaction Date & Time": "2026-09-19 14:00:00"},
            {"Transaction Date & Time": "2026-09-20 16:00:00"},
        ]

        dates = get_distinct_dates_with_counts(recs_cf, recs_eb, recs_air)
        self.assertEqual(len(dates), 3)

        # Day 1: 2026-09-18 (Friday) -> CF: 2, EB: 1, Air: 0, Total: 3
        self.assertEqual(dates[0]["date"], "2026-09-18")
        self.assertEqual(dates[0]["day_name"], "Friday")
        self.assertEqual(dates[0]["cf_count"], 2)
        self.assertEqual(dates[0]["eb_count"], 1)
        self.assertEqual(dates[0]["airtel_count"], 0)
        self.assertEqual(dates[0]["total_count"], 3)
        self.assertEqual(dates[0]["default_settlement_date"], "2026-09-19")

        # Day 2: 2026-09-19 (Saturday) -> CF: 1, EB: 0, Air: 1, Total: 2
        self.assertEqual(dates[1]["date"], "2026-09-19")
        self.assertEqual(dates[1]["day_name"], "Saturday")
        self.assertEqual(dates[1]["total_count"], 2)
        self.assertEqual(dates[1]["default_settlement_date"], "2026-09-20")

        # Day 3: 2026-09-20 (Sunday) -> CF: 0, EB: 1, Air: 1, Total: 2
        self.assertEqual(dates[2]["date"], "2026-09-20")
        self.assertEqual(dates[2]["day_name"], "Sunday")
        self.assertEqual(dates[2]["total_count"], 2)
        self.assertEqual(dates[2]["default_settlement_date"], "2026-09-21")

    def test_build_xcd_workbook_with_date_filtering(self):
        records = [
            {"SwinkPay Txn ID": "SP_D1_A", "Amount": 100.0, "Transaction Date & Time": "2026-09-18 10:00:00"},
            {"SwinkPay Txn ID": "SP_D1_B", "Amount": 200.0, "Transaction Date & Time": "2026-09-18 11:00:00"},
            {"SwinkPay Txn ID": "SP_D2_A", "Amount": 300.0, "Transaction Date & Time": "2026-09-19 12:00:00"},
            {"SwinkPay Txn ID": "SP_D3_A", "Amount": 400.0, "Transaction Date & Time": "2026-09-20 15:00:00"},
        ]

        # Case 1: Select only 1 day (Day 1: 2026-09-18)
        out_p1 = os.path.join(self.test_dir, "test_xcd_day1.xlsx")
        build_xcd_workbook(
            records=records,
            settlement_date="19-09-2026",
            output_path=out_p1,
            selected_dates=["2026-09-18"]
        )
        wb1 = openpyxl.load_workbook(out_p1)
        ws1 = wb1.active
        # Header is row 2, data rows 3 and 4 (2 records)
        self.assertEqual(ws1.cell(3, 1).value, 1)
        self.assertEqual(ws1.cell(3, 2).value, "SP_D1_A")
        self.assertEqual(ws1.cell(4, 1).value, 2)
        self.assertEqual(ws1.cell(4, 2).value, "SP_D1_B")
        self.assertEqual(ws1.cell(5, 1).value, None)  # Blank row before refund
        self.assertEqual(ws1.cell(6, 1).value, "REFUND TRANSACTIONS")

        # Case 2: Select 2 days (Day 1 & Day 2) with per-date settlement map
        out_p2 = os.path.join(self.test_dir, "test_xcd_day1_2.xlsx")
        build_xcd_workbook(
            records=records,
            settlement_date="21-09-2026",
            output_path=out_p2,
            selected_dates=["2026-09-18", "2026-09-19"],
            date_settlement_map={
                "2026-09-18": "2026-09-19",
                "2026-09-19": "2026-09-20"
            }
        )
        wb2 = openpyxl.load_workbook(out_p2)
        ws2 = wb2.active
        self.assertEqual(ws2.cell(3, 1).value, 1)
        self.assertEqual(ws2.cell(3, 2).value, "SP_D1_A")
        self.assertEqual(ws2.cell(3, 4).value, "19-09-2026")

        self.assertEqual(ws2.cell(4, 1).value, 2)
        self.assertEqual(ws2.cell(4, 2).value, "SP_D1_B")
        self.assertEqual(ws2.cell(4, 4).value, "19-09-2026")

        self.assertEqual(ws2.cell(5, 1).value, 3)
        self.assertEqual(ws2.cell(5, 2).value, "SP_D2_A")
        self.assertEqual(ws2.cell(5, 4).value, "20-09-2026")

    def test_parse_reconciliation_workbook(self):
        recon_path = "Reconciliation_2026-09-10.xlsx"
        if os.path.exists(recon_path):
            parsed = parse_reconciliation_workbook(recon_path)
            self.assertIn("stats", parsed)
            self.assertEqual(parsed["stats"]["total"]["count"], 6515)
            self.assertEqual(parsed["stats"]["cashfree"]["count"], 2691)
            self.assertEqual(parsed["stats"]["easebuzz"]["count"], 2873)
            self.assertEqual(parsed["stats"]["airtel"]["count"], 951)
            self.assertEqual(len(parsed["dates"]), 1)
            self.assertEqual(parsed["dates"][0]["date"], "2026-09-10")

    def test_upload_recon_file_and_downloads(self):
        recon_path = "Reconciliation_2026-09-10.xlsx"
        if not os.path.exists(recon_path):
            return

        with open(recon_path, "rb") as f:
            up_file = UploadFile(filename=os.path.basename(recon_path), file=f)
            resp = upload_recon_file(up_file, settlement_date="2026-09-11")

        data = json.loads(resp.body.decode("utf-8"))
        self.assertIn("session_id", data)
        self.assertEqual(data["summary"]["total_count"], 6515)
        self.assertTrue(data["xcd_status"]["all_clear"])
        self.assertIn("dates", data)

        session_id = data["session_id"]
        try:
            # Test download with dates parameter
            dl_resp = download_xcd(session_id, partner="cashfree", dates="2026-09-10")
            self.assertTrue(os.path.exists(dl_resp.path))
            self.assertEqual(dl_resp.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

            # Test download network file
            meta_path = os.path.join(SESSIONS_DIR, session_id, "session_meta.json")
            with open(meta_path, "r") as mf:
                meta = json.load(mf)
            meta["network_changes"] = [
                {
                    "SwinkPay Transaction ID": "SP_NET_TEST",
                    "Old Network": "RUPAY",
                    "New Network": "UPI_CREDIT_CARD_OFFLINE_STATIC"
                }
            ]
            with open(meta_path, "w") as mf:
                json.dump(meta, mf)

            net_resp = download_network_file(session_id)
            self.assertEqual(net_resp.filename, "Change Network.xlsx")
            self.assertTrue(os.path.exists(net_resp.path))
        finally:
            shutil.rmtree(os.path.join(SESSIONS_DIR, session_id), ignore_errors=True)

    def test_per_date_gross_and_net_amounts(self):
        records_cf = [
            {"Transaction Date & Time": "2026-09-18 10:00:00", "Amount": 100.0, "Network": "UPI"},
            {"Transaction Date & Time": "2026-09-19 11:00:00", "Amount": 200.0, "Network": "CREDIT"}
        ]
        records_eb = [
            {"Transaction Date & Time": "2026-09-18 10:15:00", "Amount": 50.0, "Network": "UPI"}
        ]
        records_air = [
            {"Transaction Date & Time": "2026-09-18 10:30:00", "Amount": 30.0, "Network": "UPI"}
        ]
        dates = get_distinct_dates_with_counts(records_cf, records_eb, records_air)
        self.assertEqual(len(dates), 2)
        d1 = dates[0]
        self.assertEqual(d1["date"], "2026-09-18")
        self.assertEqual(d1["cf_count"], 1)
        self.assertEqual(d1["cf_gross"], 100.0)
        self.assertAlmostEqual(d1["cf_net"], 98.82, places=1)
        self.assertEqual(d1["eb_count"], 1)
        self.assertEqual(d1["eb_gross"], 50.0)
        self.assertEqual(d1["airtel_count"], 1)
        self.assertEqual(d1["airtel_gross"], 30.0)
        self.assertEqual(d1["total_count"], 3)
        self.assertEqual(d1["total_gross"], 180.0)


    def test_parse_reconciliation_workbook_with_91_network_changes(self):
        recon_path = os.path.join("dashboard_sessions", "accb8824-7a6c-4e29-8d43-9f8a23062424", "Reconciliation_2026-09-22.xlsx")
        if os.path.exists(recon_path):
            parsed = parse_reconciliation_workbook(recon_path)
            changes = parsed.get("network_changes", [])
            self.assertEqual(len(changes), 91)
            from collections import Counter
            counts = Counter((c["Old Network"], c["New Network"]) for c in changes)
            self.assertEqual(counts[("RUPAY", "UPI_CREDIT_CARD_OFFLINE_STATIC")], 77)
            self.assertEqual(counts[("WA", "UPI_PPI_OFFLINE_STATIC")], 14)


if __name__ == "__main__":
    unittest.main()
