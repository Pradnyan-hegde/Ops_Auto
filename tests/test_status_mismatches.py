"""
Unit tests for Status Mismatches & Status Conflicts (PG vs CMS / SMMS).
Verifies:
1. Cashfree SUCCESS vs CMS Failed (Status 2)
2. Easebuzz 'Payment Received' vs CMS Failed
3. Cashfree FAILED vs CMS Success (Status 1)
4. Direct partner check status mismatch (Step 6)
5. Exclusion from clean settlement matched files
6. Outlet resolution with terminal mappings
7. Generation of 'Status_Mismatches' sheet in ExcelReportBuilder
"""
import os
import unittest
from unittest.mock import patch

from core.detector import ReportType
from core.reader import RawReport
from core.matcher import ReconciliationEngine
from core.aggregator import Aggregator
from core.excel_builder import ExcelReportBuilder
from core.terminal_mapper import save_merchant_terminal_mappings


def make_raw_report(report_type: ReportType, records: list, file_name: str = "report.csv", headers: list = None) -> RawReport:
    if headers is None:
        headers = list(records[0].keys()) if records else []
    raw_matrix = [headers] + [[r.get(h, "") for h in headers] for r in records]
    return RawReport(
        file_path=file_name,
        report_type=report_type,
        headers=headers,
        header_row_idx=0,
        records=records,
        raw_matrix=raw_matrix
    )


class TestStatusMismatches(unittest.TestCase):

    def setUp(self):
        self.engine = ReconciliationEngine()
        self.patcher = patch(
            "core.matcher.resolve_terminal_details",
            side_effect=lambda cand, mk=None: {
                "mms_terminal_id": "XKD8M3",
                "terminal_id": "101001189",
                "partner_ref_id": "4860",
                "branch_name": "CCD value express",
                "merchant_name": "Cafe Coffee Day",
                "vpa": ""
            } if str(cand) in ("XKD8M3", "4860", "101001189") else None
        )
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()

    def test_cf_success_vs_cms_failed_status_2(self):
        """PG shows SUCCESS, but CMS recorded as Failed (Status 2)."""
        cms_report = make_raw_report(
            report_type=ReportType.CMS,
            file_name="daily_cms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "Merchant MMS Terminal ID"],
            records=[
                {
                    "RRN/UTR": "129761680101",
                    "SwinkPay Txn ID": "SP_CF_001",
                    "Transaction Amount": "100.00",
                    "Transaction Status": "2",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 10:00:00",
                    "Merchant MMS Terminal ID": "XKD8M3"
                }
            ]
        )
        smms_report = make_raw_report(
            report_type=ReportType.SMMS,
            file_name="daily_smms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "PG/Bank", "Net Amount"],
            records=[
                {
                    "RRN/UTR": "129761680101",
                    "SwinkPay Txn ID": "SP_CF_001",
                    "Transaction Amount": "100.00",
                    "Transaction Status": "Failed",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 10:00:00",
                    "PG/Bank": "Cashfree",
                    "Net Amount": "98.50"
                }
            ]
        )
        cf_report = make_raw_report(
            report_type=ReportType.CASHFREE,
            file_name="daily_cf.csv",
            headers=["Bank Reference No.", "Reference Id", "Order Id", "Amount", "Settlement Amount", "Transaction Status", "Payment Mode", "Transaction Time"],
            records=[
                {
                    "Bank Reference No.": "129761680101",
                    "Reference Id": "SP_CF_001",
                    "Order Id": "330595-4860-ORD01",
                    "Amount": "100.00",
                    "Settlement Amount": "98.50",
                    "Transaction Status": "SUCCESS",
                    "Payment Mode": "UPI",
                    "Transaction Time": "2026-09-26 10:00:00"
                }
            ]
        )

        self.engine.set_reports(cms=cms_report, smms=smms_report, cf=cf_report)
        self.engine.run()

        # Must be recorded in status_mismatches
        self.assertEqual(len(self.engine.status_mismatches), 1)
        m = self.engine.status_mismatches[0]
        self.assertEqual(m["gateway"], "Cashfree")
        self.assertEqual(m["pg_status"], "SUCCESS")
        self.assertEqual(m["cms_status"], "Failed (Status 2)")
        self.assertEqual(m["swinkpay_txn_id"], "SP_CF_001")
        self.assertEqual(m["utr"], "129761680101")
        self.assertEqual(m["amount"], 100.0)
        self.assertIn("XKD8M3", m["outlet"])
        self.assertIn("CCD value express", m["outlet"])
        self.assertIn("SUCCESS", m["discrepancy_note"])
        self.assertIn("Failed", m["discrepancy_note"])

        # Must NOT be included in clean CMS_CF_Matched
        self.assertEqual(len(self.engine.cms_cf_matched), 0)

        # Must be added to exceptions list
        self.assertIn(m, self.engine.exceptions)

    def test_easebuzz_payment_received_vs_cms_failed(self):
        """Easebuzz shows 'Payment Received', but CMS shows 2 (Failed)."""
        cms_report = make_raw_report(
            report_type=ReportType.CMS,
            file_name="daily_cms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "Merchant MMS Terminal ID"],
            records=[
                {
                    "RRN/UTR": "129761680202",
                    "SwinkPay Txn ID": "SP_EB_002",
                    "Transaction Amount": "250.00",
                    "Transaction Status": "2",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 11:00:00",
                    "Merchant MMS Terminal ID": "XKD8M3"
                }
            ]
        )
        smms_report = make_raw_report(
            report_type=ReportType.SMMS,
            file_name="daily_smms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "PG/Bank", "Net Amount"],
            records=[
                {
                    "RRN/UTR": "129761680202",
                    "SwinkPay Txn ID": "SP_EB_002",
                    "Transaction Amount": "250.00",
                    "Transaction Status": "Failed",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 11:00:00",
                    "PG/Bank": "Easebuzz",
                    "Net Amount": "246.25"
                }
            ]
        )
        eb_report = make_raw_report(
            report_type=ReportType.EASEBUZZ,
            file_name="daily_eb.csv",
            headers=["UTR", "ID", "UPI tid", "Amount", "Status", "Payment Mode", "Transaction Date"],
            records=[
                {
                    "UTR": "129761680202",
                    "ID": "EB_9988",
                    "UPI tid": "SP_EB_002",
                    "Amount": "250.00",
                    "Status": "Payment Received",
                    "Payment Mode": "UPI",
                    "Transaction Date": "2026-09-26 11:00:00"
                }
            ]
        )

        self.engine.set_reports(cms=cms_report, smms=smms_report, eb=eb_report)
        self.engine.run()

        self.assertEqual(len(self.engine.status_mismatches), 1)
        m = self.engine.status_mismatches[0]
        self.assertEqual(m["gateway"], "Easebuzz")
        self.assertEqual(m["pg_status"], "Payment Received")
        self.assertEqual(m["cms_status"], "Failed (Status 2)")
        self.assertEqual(m["amount"], 250.0)
        self.assertEqual(len(self.engine.cms_eb_matched), 0)

    def test_cf_failed_vs_cms_success(self):
        """PG shows FAILED, but CMS recorded as Success (Status 1)."""
        cms_report = make_raw_report(
            report_type=ReportType.CMS,
            file_name="daily_cms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "Merchant MMS Terminal ID"],
            records=[
                {
                    "RRN/UTR": "129761680303",
                    "SwinkPay Txn ID": "SP_CF_003",
                    "Transaction Amount": "50.00",
                    "Transaction Status": "1",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 12:00:00",
                    "Merchant MMS Terminal ID": "XKD8M3"
                }
            ]
        )
        smms_report = make_raw_report(
            report_type=ReportType.SMMS,
            file_name="daily_smms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "PG/Bank", "Net Amount"],
            records=[
                {
                    "RRN/UTR": "129761680303",
                    "SwinkPay Txn ID": "SP_CF_003",
                    "Transaction Amount": "50.00",
                    "Transaction Status": "Success",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 12:00:00",
                    "PG/Bank": "Cashfree",
                    "Net Amount": "49.25"
                }
            ]
        )
        cf_report = make_raw_report(
            report_type=ReportType.CASHFREE,
            file_name="daily_cf.csv",
            headers=["Bank Reference No.", "Reference Id", "Order Id", "Amount", "Settlement Amount", "Transaction Status", "Payment Mode", "Transaction Time"],
            records=[
                {
                    "Bank Reference No.": "129761680303",
                    "Reference Id": "SP_CF_003",
                    "Order Id": "330595-4860-ORD03",
                    "Amount": "50.00",
                    "Settlement Amount": "0.00",
                    "Transaction Status": "FAILED",
                    "Payment Mode": "UPI",
                    "Transaction Time": "2026-09-26 12:00:00"
                }
            ]
        )

        self.engine.set_reports(cms=cms_report, smms=smms_report, cf=cf_report)
        self.engine.run()

        self.assertEqual(len(self.engine.status_mismatches), 1)
        m = self.engine.status_mismatches[0]
        self.assertEqual(m["gateway"], "Cashfree")
        self.assertEqual(m["pg_status"], "FAILED")
        self.assertEqual(m["cms_status"], "Success (Status 1)")
        self.assertEqual(len(self.engine.cms_cf_matched), 0)

    def test_direct_partner_check_status_mismatch_step_6(self):
        """Transaction in CMS (Failed) and Cashfree (SUCCESS) with SMMS Sync Status=False (not in SMMS)."""
        cms_report = make_raw_report(
            report_type=ReportType.CMS,
            file_name="daily_cms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "Merchant MMS Terminal ID", "SMMS Sync Status"],
            records=[
                {
                    "RRN/UTR": "129761680404",
                    "SwinkPay Txn ID": "SP_CF_004",
                    "Transaction Amount": "150.00",
                    "Transaction Status": "2",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 13:00:00",
                    "Merchant MMS Terminal ID": "XKD8M3",
                    "SMMS Sync Status": "false"
                }
            ]
        )
        smms_report = make_raw_report(
            report_type=ReportType.SMMS,
            file_name="daily_smms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "PG/Bank", "Net Amount"],
            records=[]
        )
        cf_report = make_raw_report(
            report_type=ReportType.CASHFREE,
            file_name="daily_cf.csv",
            headers=["Bank Reference No.", "Reference Id", "Order Id", "Amount", "Settlement Amount", "Transaction Status", "Payment Mode", "Transaction Time"],
            records=[
                {
                    "Bank Reference No.": "129761680404",
                    "Reference Id": "SP_CF_004",
                    "Order Id": "330595-4860-ORD04",
                    "Amount": "150.00",
                    "Settlement Amount": "147.75",
                    "Transaction Status": "SUCCESS",
                    "Payment Mode": "UPI",
                    "Transaction Time": "2026-09-26 13:00:00"
                }
            ]
        )

        self.engine.set_reports(cms=cms_report, smms=smms_report, cf=cf_report)
        self.engine.run()

        self.assertEqual(len(self.engine.status_mismatches), 1)
        m = self.engine.status_mismatches[0]
        self.assertEqual(m["gateway"], "Cashfree")
        self.assertEqual(m["pg_status"], "SUCCESS")
        self.assertEqual(m["cms_status"], "Failed (Status 2)")
        self.assertEqual(len(self.engine.cms_cf_matched), 0)

    def test_excel_builder_status_mismatches_sheet(self):
        """ExcelReportBuilder creates 'Status_Mismatches' sheet when conflicts exist."""
        cms_report = make_raw_report(
            report_type=ReportType.CMS,
            file_name="daily_cms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "Merchant MMS Terminal ID"],
            records=[
                {
                    "RRN/UTR": "129761680505",
                    "SwinkPay Txn ID": "SP_CF_005",
                    "Transaction Amount": "75.00",
                    "Transaction Status": "2",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 14:00:00",
                    "Merchant MMS Terminal ID": "XKD8M3"
                }
            ]
        )
        smms_report = make_raw_report(
            report_type=ReportType.SMMS,
            file_name="daily_smms.csv",
            headers=["RRN/UTR", "SwinkPay Txn ID", "Transaction Amount", "Transaction Status", "Network", "Transaction Date & Time", "PG/Bank", "Net Amount"],
            records=[
                {
                    "RRN/UTR": "129761680505",
                    "SwinkPay Txn ID": "SP_CF_005",
                    "Transaction Amount": "75.00",
                    "Transaction Status": "Failed",
                    "Network": "UPI",
                    "Transaction Date & Time": "2026-09-26 14:00:00",
                    "PG/Bank": "Cashfree",
                    "Net Amount": "73.50"
                }
            ]
        )
        cf_report = make_raw_report(
            report_type=ReportType.CASHFREE,
            file_name="daily_cf.csv",
            headers=["Bank Reference No.", "Reference Id", "Order Id", "Amount", "Settlement Amount", "Transaction Status", "Payment Mode", "Transaction Time"],
            records=[
                {
                    "Bank Reference No.": "129761680505",
                    "Reference Id": "SP_CF_005",
                    "Order Id": "330595-4860-ORD05",
                    "Amount": "75.00",
                    "Settlement Amount": "73.50",
                    "Transaction Status": "SUCCESS",
                    "Payment Mode": "UPI",
                    "Transaction Time": "2026-09-26 14:00:00"
                }
            ]
        )

        self.engine.set_reports(cms=cms_report, smms=smms_report, cf=cf_report)
        self.engine.run()

        aggregator = Aggregator(self.engine)
        aggregator.compute()

        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            builder = ExcelReportBuilder(self.engine, aggregator)
            builder.build_and_save(tmp_path)

            import openpyxl
            wb = openpyxl.load_workbook(tmp_path)
            self.assertIn("Status_Mismatches", wb.sheetnames)
            ws = wb["Status_Mismatches"]
            self.assertGreater(ws.max_row, 1)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


if __name__ == "__main__":
    unittest.main()
